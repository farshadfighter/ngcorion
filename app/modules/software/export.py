"""The fleet's software products as an Excel workbook, in English or Persian."""
from io import BytesIO
from typing import Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_FA = {
    "Software": "نرم‌افزارها",
    "Product": "محصول",
    "Identification": "شناسایی",
    "CPE": "CPE",
    "Package names": "نام بسته‌ها",
    "Publisher": "ناشر",
    "Source": "منبع",
    "Repository": "مخزن",
    "Versions": "نسخه‌ها",
    "Assets": "دارایی‌ها",
    "CVEs": "CVEها",
    "Highest severity": "بالاترین شدت",
    "Exploited (KEV)": "بهره‌برداری‌شده (KEV)",
    "Fixed in": "رفع در نسخه",
    "Yes": "بله",
    "Identified": "شناسایی‌شده",
    "Unidentified": "شناسایی‌نشده",
    "Internal": "داخلی",
    "Distribution package": "بسته توزیع",
    "Distribution": "مخزن توزیع",
    "Third-party repository": "مخزن شخص ثالث",
    "Installed manually": "نصب دستی",
    "Windows": "ویندوز",
    "Service": "سرویس",
    "Firmware": "فریم‌ور",
    "Critical": "بحرانی",
    "High": "بالا",
    "Medium": "متوسط",
    "Low": "پایین",
}

STATUS = {"known": "Identified", "unknown": "Unidentified", "internal": "Internal", "distro": "Distribution package"}
SOURCE = {"distro": "Distribution", "third_party": "Third-party repository", "manual": "Installed manually",
          "windows": "Windows", "service": "Service", "firmware": "Firmware"}
SEVERITY = {"critical": "Critical", "high": "High", "medium": "Medium", "low": "Low"}


def products_workbook(data: Dict, lang: str = "en") -> BytesIO:
    tr = (lambda s: _FA.get(s, s)) if lang == "fa" else (lambda s: s)
    wb = Workbook()
    ws = wb.active
    ws.title = tr("Software")
    ws.sheet_view.rightToLeft = lang == "fa"
    headers = ["Product", "Identification", "CPE", "Package names", "Publisher", "Source", "Repository",
               "Versions", "Assets", "CVEs", "Highest severity", "Exploited (KEV)", "Fixed in"]
    ws.append([tr(h) for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(vertical="center")
    for r in data["items"]:
        cve = r.get("cve") or {}
        versions = ", ".join(f"{v['version'] or '?'} ({v['assets']})" for v in r["versions"])
        ws.append([
            r["label"], tr(STATUS.get(r["status"], r["status"])), ", ".join(r["cpes"]), ", ".join(r["names"]),
            r.get("publisher") or "", ", ".join(tr(SOURCE.get(s, s)) for s in r["sources"]),
            ", ".join(r["origins"]), versions, r["asset_count"], cve.get("count") or "",
            tr(SEVERITY.get(cve.get("severity"), "")) if cve else "", tr("Yes") if cve.get("kev") else "",
            ", ".join(cve.get("fixed_in") or []),
        ])
    widths = [28, 16, 30, 34, 22, 22, 30, 30, 10, 8, 14, 14, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out
