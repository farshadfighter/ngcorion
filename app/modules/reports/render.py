"""
Blocks -> PDF (HTML laid out by WeasyPrint) and -> Excel (openpyxl).

The fonts ship with the module (Vazirmatn for Persian and Latin text,
DejaVu Sans Mono for IDs, IPs and versions), so the PDF looks the same on a
server with no fonts installed and without internet access.
"""
import io
import logging
import re
from pathlib import Path
from typing import Dict, List, Tuple

from jinja2 import Environment, FileSystemLoader, pass_context, select_autoescape
from markupsafe import Markup

HERE = Path(__file__).parent
FONTS = HERE / "fonts"

# WeasyPrint logs every unsupported CSS property and font-subsetting fallback;
# keep real errors only.
logging.getLogger("weasyprint").setLevel(logging.ERROR)
logging.getLogger("fontTools").setLevel(logging.ERROR)

CLASSIFICATION_COLORS = {"public": "#5d6a80", "internal": "#1e3a5f", "confidential": "#c0262d"}


def _css(value) -> Markup:
    """A value placed inside a CSS string ("...") in the <style> block."""
    text = str(value if value is not None else "")
    text = text.replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ").replace("<", "\\3c ")
    return Markup(text)


_env = Environment(loader=FileSystemLoader(str(HERE / "html")), autoescape=select_autoescape(["html"]),
                   trim_blocks=True, lstrip_blocks=True)
_env.filters["css"] = _css
_LATIN = re.compile(r"[A-Za-z0-9]")
_PERSIAN = re.compile(r"[\u0600-\u06FF]")


@pass_context
def _bidi(context, value):
    """In a Persian report, Latin text and Latin-digit values (versions, IPs, lists of
    them) keep their own direction inside right-to-left cells."""
    text = "" if value is None else str(value)
    if context.get("rtl") and _LATIN.search(text) and not _PERSIAN.search(text):
        return f"\u2066{text}\u2069"
    return text


_env.filters["bidi"] = _bidi


def _prepare(sections: List[dict]) -> List[dict]:
    """Cut long tables for the PDF and decide which table titles to print."""
    out = []
    for s in sections:
        if s.get("excel_only"):
            continue
        blocks = []
        for b in s["blocks"]:
            if b and b.get("type") == "table":
                b = dict(b)
                limit = b.get("pdf_limit")
                rows = b["rows"]
                b["shown"] = rows[:limit] if limit else rows
                b["cut"] = s["cut_note"](len(b["shown"]), len(rows)) if limit and len(rows) > limit else ""
                b["show_title"] = bool(b.get("title")) and b["title"] != s["title"]
            if b:
                blocks.append(b)
        out.append({**s, "blocks": blocks})
    return out


def _fetcher():
    """WeasyPrint may load only the bundled fonts (file: URLs inside FONTS) and
    the data: URL of the logo - never the network or another local file, even
    if a value ever reached the HTML unescaped (defence in depth: every value
    is autoescaped by the templates)."""
    from urllib.parse import unquote, urlparse

    from weasyprint.urls import URLFetcher

    fonts = FONTS.resolve()

    class LocalOnly(URLFetcher):
        """URL fetcher that serves only data: URLs and files inside the font folder."""
        def fetch(self, url, headers=None):
            """Fetch an allowed URL; anything else raises, so WeasyPrint leaves the resource out."""
            parsed = urlparse(url)
            if parsed.scheme == "data":
                return super().fetch(url, headers)
            if parsed.scheme == "file":
                path = Path(unquote(parsed.path)).resolve()
                if path.is_relative_to(fonts):
                    return super().fetch(url, headers)
            raise ValueError(f"blocked resource in a report: {url[:80]}")

    return LocalOnly(allowed_protocols=("file", "data"), allow_redirects=False)


def pdf(meta: Dict, sections: List[dict]) -> Tuple[bytes, int]:
    """Render the report as PDF and return (bytes, page count). Templates autoescape every value; fonts
    come from the bundled folder only.
    """
    from weasyprint import HTML
    html = _env.get_template("report.html").render(
        **meta, sections=_prepare(sections), classification_color=CLASSIFICATION_COLORS.get(meta["classification"],
                                                                                           "#1e3a5f"))
    document = HTML(string=html, base_url=str(FONTS) + "/", url_fetcher=_fetcher()).render()
    buf = io.BytesIO()
    document.write_pdf(buf)
    return buf.getvalue(), len(document.pages)


# ── Excel ────────────────────────────────────────────────────────────────

_BAD = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(name: str, used: set) -> str:
    base = _BAD.sub(" ", name).strip()[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _value(c):
    if isinstance(c, dict):
        v = c.get("value")
        return v if v is not None and not isinstance(v, (list, dict)) else (c.get("text") if v is None else str(v))
    return c


def _add(sheet, values) -> None:
    """Append a row; a value such as '=HYPERLINK(...)' (an asset name) stays text, never a formula."""
    from app.utils.excel_utils import append_text_safe
    append_text_safe(sheet, values)


def xlsx(meta: Dict, sections: List[dict], rtl: bool) -> bytes:
    """Render the report as an Excel workbook: one sheet per table, right-to-left for Persian; cell
    values are neutralised against formula injection.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="1E3A5F")
    head_fill = PatternFill("solid", fgColor="EEF2F7")
    wb = Workbook()
    summary = wb.active
    used = set()
    summary.title = _sheet_name(meta["t_summary"], used)
    summary.sheet_view.rightToLeft = rtl
    _add(summary, [meta["title"]])
    summary["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    for k, v in meta["about"]:
        _add(summary, [k, v])
    _add(summary, [])

    for s in sections:
        tables = [b for b in s["blocks"] if b and b.get("type") == "table"]
        other = [b for b in s["blocks"] if b and b.get("type") in ("kpis", "bullets", "note", "heading", "bars")]
        if other:
            _add(summary, [s["title"]])
            summary.cell(row=summary.max_row, column=1).font = head_font
            for b in other:
                if b["type"] == "kpis":
                    for k in b["items"]:
                        _add(summary, [k["label"], k["value"], k.get("delta") or "", k.get("sub") or ""])
                elif b["type"] == "bullets":
                    for i in b["items"]:
                        _add(summary, ["•", i])
                elif b["type"] == "bars":
                    for r in b["rows"]:
                        _add(summary, [r["label"], r["text"], r.get("change") or ""])
                elif b["type"] in ("note", "heading"):
                    _add(summary, [b["text"]])
            _add(summary, [])
        for b in tables:
            ws = wb.create_sheet(_sheet_name(b.get("sheet") or b.get("title") or s["title"], used))
            ws.sheet_view.rightToLeft = rtl
            _add(ws, [c["label"] for c in b["columns"]])
            for cell in ws[1]:
                cell.font, cell.fill = head_font, head_fill
            for row in b["rows"]:
                _add(ws, [_value(c) for c in row])
            if not b["rows"]:
                _add(ws, [b["empty"]])
            ws.freeze_panes = "A2"
            if b["rows"]:
                ws.auto_filter.ref = ws.dimensions
            for i, c in enumerate(b["columns"], start=1):
                longest = max([len(str(c["label"]))] + [len(str(_value(r[i - 1]) or "")) for r in b["rows"][:300]])
                ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, longest + 2))
                if longest > 60:
                    for cell in ws[get_column_letter(i)][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
