"""
Excel Import/Export Utilities

Handles Excel file operations for Asset and Asset Requirement data.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import List, Dict, Any
from io import BytesIO


def create_styled_workbook(title: str) -> Workbook:
    """Create a new workbook with basic styling"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    return wb


# برای تمپلیت asset list


def create_asset_list_template() -> BytesIO:
    """Creates a blank Excel template for asset list with multiple sheets."""
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    for sheet_def in ASSET_LIST_SHEETS:
        _ensure_asset_sheet(wb, sheet_def["title"], sheet_def["columns"])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def style_header_row(ws, columns: List[str]):
    """Apply styling to header row"""
    header_fill = PatternFill(
        start_color="366092", end_color="366092", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, column_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

        # Auto-width
        ws.column_dimensions[get_column_letter(col_idx)].width = max(
            len(column_name) + 2, 15
        )


ASSET_LIST_COLUMNS = [
    "ID",
    "Asset Name",
    "Hostname",
    "Asset Type",
    "Asset Role",
    "Manufacturer",
    "Model",
    "Serial Number",
    "OS Name",
    "OS Version",
    "IP Address",
    "MAC Address",
    "Ports",
    "Location",
    "Owner",
    "Status",
    "Confidentiality Level",
    "Risk Level",
    "Last Audit Date",
    "Last Patch Date",
    "Asset Value",
    "Description",
]


ASSET_REQUIREMENT_SHEETS = [
    {
        "key": "asset_types",
        "title": "Asset Types",
        "columns": ["ID", "Type Name", "Category", "Description"],
        "map_row": lambda item: [
            item.id,
            item.type_name,
            item.category,
            item.description,
        ],
    },
    {
        "key": "owners",
        "title": "Owners",
        "columns": ["ID", "Full Name", "Department", "Role", "Email", "Phone"],
        "map_row": lambda item: [
            item.id,
            item.full_name,
            item.department,
            item.role,
            item.email,
            item.phone,
        ],
    },
    {
        "key": "locations",
        "title": "Locations",
        "columns": [
            "ID",
            "Site Name",
            "Rack Name",
            "Room",
            "Floor",
            "Network Zone",
            "VLAN ID",
            "Subnet",
        ],
        "map_row": lambda item: [
            item.id,
            item.site_name,
            item.rack_name,
            item.room,
            item.floor,
            item.network_zone,
            item.vlan_id,
            item.subnet,
        ],
    },
    {
        "key": "zones",
        "title": "Network Zones",
        "columns": ["ID", "Zone Name", "Description"],
        "map_row": lambda item: [item.id, item.zone_name, item.description],
    },
    {
        "key": "os_catalog",
        "title": "OS Catalog",
        "columns": ["ID", "OS Name"],
        "map_row": lambda item: [item.id, item.os_name],
    },
    {
        "key": "vendors",
        "title": "Vendors",
        "columns": ["ID", "Vendor Name", "Vendor Type"],
        "map_row": lambda item: [item.id, item.vendor_name, item.vendor_type],
    },
    {
        "key": "dependencies",
        "title": "Dependencies",
        "columns": ["ID", "Asset ID", "Depends On ID", "Relation Type", "Description"],
        "map_row": lambda item: [
            item.id,
            item.asset_id,
            item.depends_on_id,
            item.relation_type.value
            if hasattr(item, "relation_type") and item.relation_type
            else "",
            item.description,
        ],
    },
]


ASSET_LIST_SHEETS = [
    {
        "key": "overview",
        "title": "Overview",
        "columns": ["Asset Name", "Hostname", "Type", "Role", "Manufacturer", "Model"],
        "map_row": lambda asset: [
            asset.asset_name,
            asset.hostname,
            asset.asset_type.type_name if asset.asset_type else "",  # اصلاح شد: .type_name به جای .name
            asset.asset_role if hasattr(asset, 'asset_role') else getattr(asset, 'role', ''), # معمولا asset_role است
            asset.manufacturer if hasattr(asset, 'manufacturer') else getattr(asset, 'vendor', ''),
            asset.model,
        ],
    },
    {
        "key": "network_system",
        "title": "Network & System",
        "columns": ["Asset Name", "Serial", "OS", "IP Address", "MAC Address", "Ports"],
        "map_row": lambda asset: [
            asset.asset_name,
            getattr(asset, 'serial_number', getattr(asset, 'serial', '')), 
            asset.os_name if hasattr(asset, 'os_name') else (f"{asset.os.os_name} {asset.os_version}" if hasattr(asset, 'os') and asset.os else ""), # اصلاح شد
            asset.ip_address,
            asset.mac_address,
            getattr(asset, "ports", "N/A"),
        ],
    },
    {
        "key": "location_owner",
        "title": "Location & Owner",
        "columns": ["Asset Name", "Location", "Owner", "Status"],
        "map_row": lambda asset: [
            asset.asset_name,
            asset.location.site_name if asset.location else "",  # اصلاح شد: .site_name به جای .name
            asset.owner.full_name if asset.owner else "",        # اصلاح شد: .full_name به جای .name
            asset.status.value if hasattr(asset.status, 'value') else asset.status, # برای enum
        ],
    },
    {
        "key": "security_audit",
        "title": "Security & Audit",
        "columns": [
            "Asset Name",
            "Confidentiality Level",
            "Risk Level",
            "Last Audit Date",
            "Last Patch Date",
        ],
        "map_row": lambda asset: [
            asset.asset_name,
            asset.confidentiality_level.value if hasattr(asset, 'confidentiality_level') and hasattr(asset.confidentiality_level, 'value') else getattr(asset, 'confidentiality', ''),
            asset.risk_level.value if hasattr(asset.risk_level, 'value') else asset.risk_level,
            asset.last_audit_date.strftime("%Y-%m-%d") if asset.last_audit_date else "",
            asset.last_patch_date.strftime("%Y-%m-%d") if asset.last_patch_date else "",
        ],
    },
]



def _build_asset_row(asset: Any) -> List[Any]:
    ports_str = (
        ", ".join([str(p.port_number) for p in asset.ports]) if asset.ports else "N/A"
    )
    return [
        asset.id,
        asset.asset_name,
        asset.hostname,
        asset.asset_type.type_name if asset.asset_type else "",
        asset.asset_role,
        asset.manufacturer,
        asset.model,
        asset.serial_number,
        asset.os_name,
        asset.os_version,
        asset.ip_address,
        asset.mac_address,
        ports_str,
        asset.location.site_name if asset.location else "",
        asset.owner.full_name if asset.owner else "",
        asset.status.value if asset.status else "",
        asset.confidentiality_level.value if asset.confidentiality_level else "",
        asset.risk_level.value if asset.risk_level else "",
        asset.last_audit_date.strftime("%Y-%m-%d") if asset.last_audit_date else "",
        asset.last_patch_date.strftime("%Y-%m-%d") if asset.last_patch_date else "",
        float(asset.asset_value) if asset.asset_value is not None else "",
        asset.description,
    ]


def _ensure_asset_sheet(wb: Workbook, title: str, columns: List[str]):
    ws = wb.create_sheet(title)
    style_header_row(ws, columns)
    return ws


"""
    def export_assets_to_excel(assets: List[Any], include_data: bool = True) -> BytesIO:
        '''
        Export assets to Excel file

        Args:
            assets: List of Asset objects
            include_data: If False, creates empty template. If True, includes asset data.

        Returns:
            BytesIO object containing the Excel file
        '''
        wb = create_styled_workbook("Assets")
        ws = wb.active

        # Style header
        style_header_row(ws, ASSET_LIST_COLUMNS)

        # Add data if requested
        if include_data and assets:
            for row_idx, asset in enumerate(assets, 2):
                for col_idx, value in enumerate(_build_asset_row(asset), start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value)

        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
"""


# export with sheet (new)
def export_assets_to_excel(assets: List[Any], include_data: bool = True) -> BytesIO:
    """
    Export assets to Excel file

    Args:
        assets: List of Asset objects
        include_data: If False, creates empty template. If True, includes asset data.

    Returns:
        BytesIO object containing the Excel file
    """
    if not include_data:
        return create_asset_list_template()

    return export_assets_to_excel_grouped(assets)


# new export for asset list with sheet


def export_assets_to_excel_grouped(assets: List[Any]) -> BytesIO:
    """Exports a list of assets to a multi-sheet Excel file."""
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    for sheet_def in ASSET_LIST_SHEETS:
        ws = _ensure_asset_sheet(wb, sheet_def["title"], sheet_def["columns"])
        map_row_func = sheet_def["map_row"]
        for asset in assets:
            row_data = map_row_func(asset)
            ws.append(row_data)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def create_asset_template() -> BytesIO:
    """
    Create an empty Excel template for asset import

    Returns:
        BytesIO object containing the template Excel file
    """
    # return export_assets_to_excel([], include_data=False)
    return create_asset_list_template()


def import_assets_from_excel(
    file_content: BytesIO, db_session, current_user
) -> Dict[str, Any]:
    """
    Import assets from an Excel file

    Args:
        file_content: BytesIO object containing the Excel file
        db_session: Database session
        current_user: Current authenticated user

    Returns:
        Dictionary with import results:
        - created: Number of assets created
        - updated: Number of assets updated
        - errors: List of error messages
        - details: Detailed results for each row
    """
    from openpyxl import load_workbook
    from app.models import Asset, AssetType, AssetLocation, AssetOwner
    from app.models.enums import StatusEnum, ConfidentialityLevelEnum, RiskLevelEnum
    from datetime import datetime

    results = {"created": 0, "updated": 0, "skipped": 0, "errors": [], "details": []}

    try:
        # Load workbook
        wb = load_workbook(file_content, data_only=True)
        ws = wb.active

        # Get header row
        headers = [cell.value for cell in ws[1]]

        # Map headers to field names
        header_map = {
            "ID": "id",
            "Asset Name": "asset_name",
            "Hostname": "hostname",
            "Asset Type": "asset_type",
            "Asset Role": "asset_role",
            "Manufacturer": "manufacturer",
            "Model": "model",
            "Serial Number": "serial_number",
            "OS Name": "os_name",
            "OS Version": "os_version",
            "IP Address": "ip_address",
            "MAC Address": "mac_address",
            "Location": "location",
            "Owner": "owner",
            "Status": "status",
            "Confidentiality Level": "confidentiality_level",
            "Risk Level": "risk_level",
            "Last Audit Date": "last_audit_date",
            "Last Patch Date": "last_patch_date",
            "Asset Value": "asset_value",
            "Description": "description",
        }

        # Process each row (skip header)
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            try:
                # Build data dictionary
                row_data = {}
                for col_idx, value in enumerate(row):
                    if col_idx < len(headers) and headers[col_idx] in header_map:
                        field_name = header_map[headers[col_idx]]
                        if value is not None and value != "":
                            row_data[field_name] = value

                # Skip empty rows
                if not row_data or not row_data.get("asset_name"):
                    results["skipped"] += 1
                    continue

                # Resolve foreign keys
                asset_type_id = None
                if "asset_type" in row_data:
                    asset_type = (
                        db_session.query(AssetType)
                        .filter(AssetType.type_name == row_data["asset_type"])
                        .first()
                    )
                    if asset_type:
                        asset_type_id = asset_type.id
                    else:
                        results["errors"].append(
                            f"Row {row_idx}: Asset type '{row_data['asset_type']}' not found"
                        )
                        continue
                    del row_data["asset_type"]

                location_id = None
                if "location" in row_data:
                    location = (
                        db_session.query(AssetLocation)
                        .filter(AssetLocation.site_name == row_data["location"])
                        .first()
                    )
                    if location:
                        location_id = location.id
                    del row_data["location"]

                owner_id = None
                if "owner" in row_data:
                    owner = (
                        db_session.query(AssetOwner)
                        .filter(AssetOwner.full_name == row_data["owner"])
                        .first()
                    )
                    if owner:
                        owner_id = owner.id
                    del row_data["owner"]

                # Convert enums
                if "status" in row_data:
                    try:
                        row_data["status"] = StatusEnum(row_data["status"])
                    except ValueError:
                        results["errors"].append(
                            f"Row {row_idx}: Invalid status '{row_data['status']}'"
                        )
                        del row_data["status"]

                if "confidentiality_level" in row_data:
                    try:
                        row_data["confidentiality_level"] = ConfidentialityLevelEnum(
                            row_data["confidentiality_level"]
                        )
                    except ValueError:
                        del row_data["confidentiality_level"]

                if "risk_level" in row_data:
                    try:
                        row_data["risk_level"] = RiskLevelEnum(row_data["risk_level"])
                    except ValueError:
                        del row_data["risk_level"]

                # Convert dates
                for date_field in ["last_audit_date", "last_patch_date"]:
                    if date_field in row_data and isinstance(row_data[date_field], str):
                        try:
                            row_data[date_field] = datetime.strptime(
                                row_data[date_field], "%Y-%m-%d"
                            ).date()
                        except ValueError:
                            del row_data[date_field]

                # Add foreign keys
                if asset_type_id:
                    row_data["asset_type_id"] = asset_type_id
                if location_id:
                    row_data["location_id"] = location_id
                if owner_id:
                    row_data["owner_id"] = owner_id

                # Check if asset exists (by ID, asset_name, or IP)
                existing_asset = None
                if "id" in row_data and row_data["id"]:
                    existing_asset = (
                        db_session.query(Asset)
                        .filter(Asset.id == row_data["id"])
                        .first()
                    )
                elif "ip_address" in row_data:
                    existing_asset = (
                        db_session.query(Asset)
                        .filter(
                            Asset.ip_address == row_data["ip_address"],
                            Asset.user_id == current_user.id,
                        )
                        .first()
                    )

                if not existing_asset and "asset_name" in row_data:
                    existing_asset = (
                        db_session.query(Asset)
                        .filter(
                            Asset.asset_name == row_data["asset_name"],
                            Asset.user_id == current_user.id,
                        )
                        .first()
                    )

                # Remove ID from row_data for new assets
                if "id" in row_data and not existing_asset:
                    del row_data["id"]

                if existing_asset:
                    # Update existing asset
                    if "id" in row_data:
                        del row_data["id"]  # Don't update ID
                    for key, value in row_data.items():
                        if hasattr(existing_asset, key):
                            setattr(existing_asset, key, value)

                    results["updated"] += 1
                    results["details"].append(
                        {
                            "row": row_idx,
                            "action": "updated",
                            "asset_name": row_data.get(
                                "asset_name", existing_asset.asset_name
                            ),
                            "id": existing_asset.id,
                        }
                    )
                else:
                    # Create new asset
                    if not asset_type_id:
                        results["errors"].append(
                            f"Row {row_idx}: Asset type is required for new assets"
                        )
                        continue

                    row_data["user_id"] = current_user.id
                    new_asset = Asset(**row_data)
                    db_session.add(new_asset)
                    db_session.flush()  # Get ID before commit

                    results["created"] += 1
                    results["details"].append(
                        {
                            "row": row_idx,
                            "action": "created",
                            "asset_name": row_data.get("asset_name"),
                            "id": new_asset.id,
                        }
                    )

            except Exception as e:
                results["errors"].append(f"Row {row_idx}: {str(e)}")
                continue

        # Commit all changes
        db_session.commit()

    except Exception as e:
        db_session.rollback()
        results["errors"].append(f"Fatal error: {str(e)}")

    return results


# ====================================
# Asset Requirements Export Functions
# ====================================


def export_asset_requirements_to_excel(data_dict: Dict[str, List[Any]]) -> BytesIO:
    """
    Export all asset requirements to a single Excel file with multiple sheets

    Args:
        data_dict: Dictionary with sheet names as keys and data lists as values
            Expected keys: asset_types, owners, locations, zones, os_catalog, vendors, dependencies

    Returns:
        BytesIO object containing the Excel file
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet in ASSET_REQUIREMENT_SHEETS:
        ws = _ensure_asset_sheet(wb, sheet["title"], sheet["columns"])
        rows = data_dict.get(sheet["key"]) or []

        for row_idx, item in enumerate(rows, start=2):
            row_values = sheet["map_row"](item)
            for col_idx, value in enumerate(row_values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def create_asset_requirements_template() -> BytesIO:
    """
    Create an empty Excel template for asset requirements import

    Returns:
        BytesIO object containing the template Excel file with all sheets
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet in ASSET_REQUIREMENT_SHEETS:
        _ensure_asset_sheet(wb, sheet["title"], sheet["columns"])

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_asset_requirements_from_excel(
    file_content: BytesIO, db_session, current_user
) -> Dict[str, Any]:
    """
    Import asset requirements from an Excel file with multiple sheets

    Args:
        file_content: BytesIO object containing the Excel file
        db_session: Database session
        current_user: Current authenticated user

    Returns:
        Dictionary with import results for each sheet
    """
    from openpyxl import load_workbook
    from app.models import (
        AssetType,
        AssetOwner,
        AssetLocation,
        NetworkZone,
        OSCatalog,
        VendorCatalog,
        AssetDependency,
    )
    from app.models.enums import RelationTypeEnum

    results = {
        "asset_types": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "owners": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "locations": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "zones": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "os_catalog": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "vendors": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
        "dependencies": {"created": 0, "updated": 0, "skipped": 0, "errors": []},
    }

    try:
        wb = load_workbook(file_content, data_only=True)

        # Process Asset Types sheet
        if "Asset Types" in wb.sheetnames:
            ws = wb["Asset Types"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if Type Name is empty
                        results["asset_types"]["skipped"] += 1
                        continue

                    data = {
                        "type_name": row[1],
                        "category": row[2],
                        "description": row[3] if len(row) > 3 else None,
                    }

                    # Check if exists by ID or name
                    existing = None
                    if row[0]:  # ID provided
                        existing = (
                            db_session.query(AssetType)
                            .filter(AssetType.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(AssetType)
                            .filter(AssetType.type_name == data["type_name"])
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None:
                                setattr(existing, key, value)
                        results["asset_types"]["updated"] += 1
                    else:
                        new_item = AssetType(**data)
                        db_session.add(new_item)
                        results["asset_types"]["created"] += 1

                except Exception as e:
                    results["asset_types"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process Owners sheet
        if "Owners" in wb.sheetnames:
            ws = wb["Owners"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if Full Name is empty
                        results["owners"]["skipped"] += 1
                        continue

                    data = {
                        "full_name": row[1],
                        "department": row[2] if len(row) > 2 else None,
                        "role": row[3] if len(row) > 3 else None,
                        "email": row[4] if len(row) > 4 else None,
                        "phone": row[5] if len(row) > 5 else None,
                        "user_id": current_user.id,
                    }

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(AssetOwner)
                            .filter(AssetOwner.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(AssetOwner)
                            .filter(
                                AssetOwner.full_name == data["full_name"],
                                AssetOwner.user_id == current_user.id,
                            )
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None and key != "user_id":
                                setattr(existing, key, value)
                        results["owners"]["updated"] += 1
                    else:
                        new_item = AssetOwner(**data)
                        db_session.add(new_item)
                        results["owners"]["created"] += 1

                except Exception as e:
                    results["owners"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process Locations sheet
        if "Locations" in wb.sheetnames:
            ws = wb["Locations"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if Site Name is empty
                        results["locations"]["skipped"] += 1
                        continue

                    data = {
                        "site_name": row[1],
                        "rack_name": row[2] if len(row) > 2 else None,
                        "room": row[3] if len(row) > 3 else None,
                        "floor": row[4] if len(row) > 4 else None,
                        "network_zone": row[5] if len(row) > 5 else None,
                        "vlan_id": row[6] if len(row) > 6 else None,
                        "subnet": row[7] if len(row) > 7 else None,
                        "user_id": current_user.id,
                    }

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(AssetLocation)
                            .filter(AssetLocation.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(AssetLocation)
                            .filter(
                                AssetLocation.site_name == data["site_name"],
                                AssetLocation.user_id == current_user.id,
                            )
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None and key != "user_id":
                                setattr(existing, key, value)
                        results["locations"]["updated"] += 1
                    else:
                        new_item = AssetLocation(**data)
                        db_session.add(new_item)
                        results["locations"]["created"] += 1

                except Exception as e:
                    results["locations"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process Network Zones sheet
        if "Network Zones" in wb.sheetnames:
            ws = wb["Network Zones"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if Zone Name is empty
                        results["zones"]["skipped"] += 1
                        continue

                    data = {
                        "zone_name": row[1],
                        "description": row[2] if len(row) > 2 else None,
                    }

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(NetworkZone)
                            .filter(NetworkZone.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(NetworkZone)
                            .filter(NetworkZone.zone_name == data["zone_name"])
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None:
                                setattr(existing, key, value)
                        results["zones"]["updated"] += 1
                    else:
                        new_item = NetworkZone(**data)
                        db_session.add(new_item)
                        results["zones"]["created"] += 1

                except Exception as e:
                    results["zones"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process OS Catalog sheet
        if "OS Catalog" in wb.sheetnames:
            ws = wb["OS Catalog"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if OS Name is empty
                        results["os_catalog"]["skipped"] += 1
                        continue

                    data = {"os_name": row[1]}

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(OSCatalog)
                            .filter(OSCatalog.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(OSCatalog)
                            .filter(OSCatalog.os_name == data["os_name"])
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None:
                                setattr(existing, key, value)
                        results["os_catalog"]["updated"] += 1
                    else:
                        new_item = OSCatalog(**data)
                        db_session.add(new_item)
                        results["os_catalog"]["created"] += 1

                except Exception as e:
                    results["os_catalog"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process Vendors sheet
        if "Vendors" in wb.sheetnames:
            ws = wb["Vendors"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if not row[1]:  # Skip if Vendor Name is empty
                        results["vendors"]["skipped"] += 1
                        continue

                    data = {
                        "vendor_name": row[1],
                        "vendor_type": row[2] if len(row) > 2 else None,
                    }

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(VendorCatalog)
                            .filter(VendorCatalog.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(VendorCatalog)
                            .filter(VendorCatalog.vendor_name == data["vendor_name"])
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None:
                                setattr(existing, key, value)
                        results["vendors"]["updated"] += 1
                    else:
                        new_item = VendorCatalog(**data)
                        db_session.add(new_item)
                        results["vendors"]["created"] += 1

                except Exception as e:
                    results["vendors"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Process Dependencies sheet
        if "Dependencies" in wb.sheetnames:
            ws = wb["Dependencies"]
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                try:
                    if len(row) < 2 or not row[1] or not row[2]:
                        results["dependencies"]["skipped"] += 1
                        continue

                    relation_type = row[3] if len(row) > 3 else None
                    if relation_type is None:
                        results["dependencies"]["errors"].append(
                            f"Row {row_idx}: Relation Type is required"
                        )
                        continue
                    relation_type = str(relation_type).strip().lower().replace(" ", "_")

                    relation_type_value = None
                    try:
                        relation_type_value = RelationTypeEnum(relation_type)
                    except ValueError:
                        results["dependencies"]["errors"].append(
                            f"Row {row_idx}: Invalid relation type '{relation_type}'"
                        )
                        continue

                    data = {
                        "asset_id": int(row[1]),
                        "depends_on_id": int(row[2]),
                        "relation_type": relation_type_value,
                        "description": row[4] if len(row) > 4 else None,
                    }

                    existing = None
                    if row[0]:
                        existing = (
                            db_session.query(AssetDependency)
                            .filter(AssetDependency.id == row[0])
                            .first()
                        )
                    if not existing:
                        existing = (
                            db_session.query(AssetDependency)
                            .filter(
                                AssetDependency.asset_id == data["asset_id"],
                                AssetDependency.depends_on_id == data["depends_on_id"],
                            )
                            .first()
                        )

                    if existing:
                        for key, value in data.items():
                            if value is not None:
                                setattr(existing, key, value)
                        results["dependencies"]["updated"] += 1
                    else:
                        new_item = AssetDependency(**data)
                        db_session.add(new_item)
                        results["dependencies"]["created"] += 1

                except Exception as e:
                    results["dependencies"]["errors"].append(f"Row {row_idx}: {str(e)}")

        # Commit all changes
        db_session.commit()

    except Exception as e:
        db_session.rollback()
        for sheet in results:
            results[sheet]["errors"].append(f"Fatal error: {str(e)}")

    return results
