"""
Software inventory (app/modules/software): parsing what each platform
reports, saving collections and their changes, identifying products, the
CVE findings that come from installed software, the API and the alert rule.

The Debian sample in tests/mock_data/software is real output of the three
collection commands on an Ubuntu 24.04 host with Docker's repository and the
ondrej/php and deadsnakes PPAs, cut down to a few packages.
"""
import asyncio
import importlib
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.core.database import engine as db_engine
from app.models import Asset, User, UserRole
from app.models.asset_log import AssetLog
from app.models.asset_types import AssetType
from app.models.software import SoftwareChange, SoftwareCollection, SoftwareItem, SoftwareProductMap
from app.modules.alerts import events as alert_events
from app.modules.cve import cpe as cve_cpe
from app.modules.cve import feeds, findings
from app.modules.cve import settings as cve_settings
from app.modules.cve import store as cve_store
from app.modules.software import catalog, collect, hooks, identify, service, store

sys.path.insert(0, str(Path(__file__).parent))
from test_cve import FORTI, _match, _nvd_cve, nvd_page  # noqa: E402

api = importlib.import_module("app.modules.software.router")
DATA = Path(__file__).parent / "mock_data" / "software"


def _debian_outputs():
    return {k: (DATA / f"ubuntu2404_{k}.txt").read_text() for k in ("deb_packages", "deb_sources", "deb_repos")}


@pytest.fixture
def db():
    connection = db_engine.connect()
    trans = connection.begin()
    session = Session(bind=connection, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        session.close()
        trans.rollback()
        connection.close()


_seq = [0]


def _n():
    _seq[0] += 1
    return _seq[0]


def _user(db, role=UserRole.ADMIN):
    n = _n()
    u = User(username=f"sw_user_{n}", email=f"sw{n}@example.com", hashed_password="x", role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _asset(db, **kw):
    t = AssetType(type_name=f"Server {_n()}", category="server")
    db.add(t)
    db.flush()
    a = Asset(asset_name=f"sw-asset-{_n()}", asset_type_id=t.id, ip_address=f"10.88.{_n() % 250}.{_n() % 250}", **kw)
    db.add(a)
    db.flush()
    return a


def _pkg(name, version, source="third_party", origin="download.docker.com", **kw):
    return {"name": name, "version": version, "arch": "amd64", "source": source, "origin": origin,
            "publisher": None, "kind": "package", "source_package": kw.pop("source_package", name), "vkind": "deb",
            **kw}


def _by_name(items):
    return {i["name"]: i for i in items}


# ── parsing ───────────────────────────────────────────────────────────────

def test_debian_sources_from_real_apt_output():
    items = _by_name(collect.parse_debian(_debian_outputs()))
    # dpkg's "rc" (removed, config kept) is not installed
    assert "old-removed" not in items
    assert items["docker-ce"]["source"] == "third_party" and items["docker-ce"]["origin"] == "download.docker.com"
    assert items["php8.4-cli"]["origin"] == "PPA ondrej-php" and items["php8.4-cli"]["source_package"] == "php8.4"
    assert items["python3.13"]["origin"] == "PPA deadsnakes"
    assert items["openssh-server"]["source"] == "distro" and items["openssh-server"]["origin"] == "Ubuntu"
    # installed version no longer in any index, but its other versions are Ubuntu's: still a distro package
    assert items["libssl3t64"]["source"] == "distro"
    # the dpkg architecture suffix is not part of the name
    assert "libssl3t64:amd64" not in items and items["libssl3t64"]["arch"] == "amd64"
    # a .deb installed by hand: no repository knows it
    assert items["acme-monitor"]["source"] == "manual" and items["acme-monitor"]["origin"] is None


def test_rhel_repositories_vendor_and_source_rpm():
    outputs = {
        "rpm_packages": "\n".join([
            "openssl\t1:3.0.7-27.el9\tx86_64\tRocky Enterprise Software Foundation\topenssl-3.0.7-27.el9.src.rpm",
            "nginx\t1:1.26.1-1.el9.ngx\tx86_64\tNginx, Inc.\tnginx-1.26.1-1.el9.ngx.src.rpm",
            "htop\t0:3.3.0-1.el9\tx86_64\tFedora Project\thtop-3.3.0-1.el9.src.rpm",
            "acme-agent\t0:1.0-1\tx86_64\t(none)\tacme-agent-1.0-1.src.rpm",
            "gpg-pubkey\t0:350d275d-6279464b\t(none)\t(none)\t(none)",
        ]),
        "rpm_repos": "openssl\tx86_64\tbaseos\nnginx\tx86_64\tnginx-stable\nhtop\tx86_64\tepel\n"
                     "acme-agent\tx86_64\t@commandline\n",
    }
    items = _by_name(collect.parse_rhel(outputs))
    assert "gpg-pubkey" not in items
    assert items["openssl"]["source"] == "distro" and items["openssl"]["version"] == "1:3.0.7-27.el9"
    assert items["nginx"]["source"] == "third_party" and items["nginx"]["origin"] == "nginx-stable"
    assert items["htop"]["origin"] == "EPEL"
    assert items["acme-agent"]["source"] == "manual"
    assert items["htop"]["version"] == "3.3.0-1.el9"         # epoch 0 dropped
    assert items["nginx"]["source_package"] == "nginx"


def test_windows_programs_updates_and_build():
    out = "\r\n".join([
        "P\tGoogle Chrome\t118.0.5993.118\tGoogle LLC\t20231020\tx64",
        "P\t7-Zip 22.01 (x64)\t22.01\tIgor Pavlov\t\tx64",
        "P\t7-Zip 22.01 (x64)\t22.01\tIgor Pavlov\t\tx64",
        "H\tKB5031364\tSecurity Update\t20231011",
        "B\tWindows Server 2022 Datacenter\t10.0.20348.2031\t21H2",
        "garbage line",
    ])
    items = collect.parse_windows(out)
    kinds = [(i["kind"], i["name"]) for i in items]
    assert kinds == [("program", "Google Chrome"), ("program", "7-Zip 22.01 (x64)"), ("hotfix", "KB5031364"),
                     ("os", "Windows Server 2022 Datacenter")]
    assert items[0]["publisher"] == "Google LLC" and items[0]["source"] == "windows"


def test_single_versions_from_services_and_devices():
    assert collect.apache_item("Server version: Apache/2.4.52 (Ubuntu)")["source"] == "distro"
    plain = collect.apache_item("Server version: Apache/2.4.58 (Unix)")
    assert plain["source"] == "service" and plain["version"] == "2.4.58"
    assert collect.apache_item("command not found") is None
    assert collect.mongodb_item("db version v6.0.12\ngit version: abc")["version"] == "6.0.12"
    sql = collect.mssql_item("Microsoft SQL Server 2019 (RTM-CU22) (KB5027702) - 15.0.4322.2 (X64)")
    assert sql["name"] == "Microsoft SQL Server 2019" and sql["version"] == "15.0.4322.2"
    ios = collect.cisco_item("Cisco IOS XE Software, Version 17.09.04a\nCisco IOS Software [Cupertino], ...")
    assert ios["name"] == "Cisco IOS XE" and ios["version"] == "17.09.04a" and ios["kind"] == "firmware"
    fgt = collect.fortinet_item({"fortios_version": "7.2.5", "model": "FGT60F", "build": "1517"})
    assert fgt["name"] == "Fortinet FortiOS" and fgt["source"] == "firmware"
    assert collect.fortinet_item({"fortios_version": "0.0.0"}) is None


def test_catalog_keys_and_upstream_versions():
    assert catalog.match_key("libssl3t64:amd64") == "libssl3t64"
    assert catalog.match_key("Mozilla Firefox (x64 en-US)", windows=True) == catalog.match_key("Mozilla Firefox",
                                                                                                 windows=True)
    assert catalog.lookup("docker-ce")[1][0] == ("docker", "docker")
    assert catalog.lookup("openssh-server")[1] == [("openbsd", "openssh")]
    assert catalog.lookup("postgresql-16")[1] == [("postgresql", "postgresql")]
    assert catalog.is_component("libfoo-dev") and not catalog.is_component("nginx")
    assert catalog.upstream_version("5:29.3.1-1~ubuntu.24.04~noble", "deb") == "29.3.1"
    assert catalog.upstream_version("1:9.6p1-3ubuntu13.19", "deb") == "9.6p1"
    assert catalog.upstream_version("1:1.26.1-1.el9.ngx", "rpm") == "1.26.1"
    assert catalog.catalog_size() > 100


def test_identify_known_unknown_distro_and_admin_answer():
    items = _by_name(collect.parse_debian(_debian_outputs()))
    maps = {}
    docker = identify.resolve(items["docker-ce"], maps)
    assert docker.status == "known" and docker.nvd and docker.version == "29.3.1"
    php = identify.resolve(items["php8.4-cli"], maps)
    assert php.status == "known" and php.label == "PHP"     # through its source package php8.4
    ssh = identify.resolve(items["openssh-server"], maps)
    assert ssh.status == "known" and not ssh.nvd           # distro: waits for the distribution's advisories
    assert identify.resolve(items["coreutils"], maps).status == "distro"
    assert identify.resolve(items["acme-monitor"], maps).status == "unknown"

    maps = {"acme-monitor": SoftwareProductMap(match_key="acme-monitor", status="cpe", vendor="acme",
                                               product="monitor")}
    acme = identify.resolve(items["acme-monitor"], maps)
    assert acme.status == "known" and acme.cpes == [("acme", "monitor")] and acme.mapped_by == "admin"
    maps = {"acme-monitor": SoftwareProductMap(match_key="acme-monitor", status="internal")}
    assert identify.resolve(items["acme-monitor"], maps).status == "internal"
    # docker-ce and docker-ce-cli are one product for NVD
    ids = identify.nvd_identities(list(items.values()), {})
    assert sum(1 for v, p, *_ in ids if (v, p) == ("docker", "docker")) == 1
    assert not any(p == "openssh" for _, p, *_ in ids)


# ── saving ────────────────────────────────────────────────────────────────

def test_first_collection_has_no_changes_then_diffs_are_recorded(db):
    asset = _asset(db)
    first = store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7"), _pkg("telnet", "0.17", source="distro"),
                                               _pkg("nginx", "1.24.0", origin="nginx.org")])
    assert first.status == "ok" and first.summary["first"] is True
    assert db.query(SoftwareChange).filter_by(asset_id=asset.id).count() == 0

    second = store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.9"), _pkg("nginx", "1.24.0", origin="nginx.org"),
                                                _pkg("acme-monitor", "2.3.1", source="manual", origin=None)])
    assert (second.added, second.updated, second.removed) == (1, 1, 1)
    changes = {(c.change, c.name) for c in db.query(SoftwareChange).filter_by(collection_id=second.id)}
    assert changes == {("added", "acme-monitor"), ("updated", "docker-ce"), ("removed", "telnet")}
    names = {i.name: i for i in db.query(SoftwareItem).filter_by(asset_id=asset.id)}
    assert set(names) == {"docker-ce", "nginx", "acme-monitor"} and names["docker-ce"].version == "24.0.9"
    assert second.summary["by_source"] == {"third_party": 2, "manual": 1}


def test_failed_collection_keeps_the_previous_list_and_other_collectors_are_untouched(db):
    asset = _asset(db)
    store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7")])
    store.save(db, asset.id, "apache", [collect.apache_item("Server version: Apache/2.4.58 (Unix)")])
    failed = store.save(db, asset.id, "linux", None, error="The package list could not be read")
    assert failed.status == "failed"
    assert {i.name for i in db.query(SoftwareItem).filter_by(asset_id=asset.id)} == {"docker-ce", "Apache HTTP Server"}
    store.save(db, asset.id, "linux", [_pkg("nginx", "1.26.0", origin="nginx.org")])
    assert {i.name for i in db.query(SoftwareItem).filter_by(asset_id=asset.id)} == {"nginx", "Apache HTTP Server"}


def test_changes_older_than_a_year_are_purged(db):
    asset = _asset(db)
    col = store.save(db, asset.id, "linux", [_pkg("a", "1")])
    db.add(SoftwareChange(asset_id=asset.id, collection_id=col.id, change="added", name="old",
                          at=datetime.utcnow() - timedelta(days=400)))
    db.flush()
    store.save(db, asset.id, "linux", [_pkg("a", "2")])
    assert {c.name for c in db.query(SoftwareChange).filter_by(asset_id=asset.id)} == {"a"}


def test_firmware_version_updates_the_asset_and_its_history(db):
    user = _user(db)
    asset = _asset(db, os_name="FortiOS", os_version="7.2.4")
    hooks.save_single(db, asset.id, "fortinet", collect.fortinet_item({"fortios_version": "7.2.5"}), user_id=user.id)
    db.refresh(asset)
    assert asset.os_version == "7.2.5"
    log = db.query(AssetLog).filter_by(asset_id=asset.id).order_by(AssetLog.id.desc()).first()
    assert log is not None and "7.2.5" in str(log.details)
    # the same version again changes nothing
    hooks.save_single(db, asset.id, "fortinet", collect.fortinet_item({"fortios_version": "7.2.5"}), user_id=user.id)
    assert db.query(AssetLog).filter_by(asset_id=asset.id).count() == 1
    # nothing read: nothing recorded
    assert hooks.save_single(db, asset.id, "cisco", None) is None


def test_linux_hook_records_a_failed_collection_without_raising(db):
    asset = _asset(db)

    class Broken:
        def collect_audit_data(self, cmds):
            raise TimeoutError("socket timed out")

    raw = hooks.collect_linux(Broken(), "ubuntu")
    assert raw["family"] == "debian" and "TimeoutError" in raw["error"]
    col = hooks.save_linux(db, asset.id, raw)
    assert col.status == "failed"

    class Fake:
        def collect_audit_data(self, cmds):
            assert all(c["sudo"] is False for c in cmds)       # read-only, never sudo
            return _debian_outputs()

    col = hooks.save_linux(db, asset.id, hooks.collect_linux(Fake(), "ubuntu"))
    assert col.status == "ok" and col.item_count == 24


# ── CVE findings from installed software ─────────────────────────────────

def _load_cves(db):
    cve_store.apply_cves(db, feeds.parse_nvd_page(nvd_page(
        _nvd_cve("CVE-2024-41110", score=9.9, matches=[
            _match("cpe:2.3:a:docker:docker:*:*:*:*:*:*:*:*", start_incl="19.03.0", end_excl="27.1.1")]),
        _nvd_cve("CVE-2023-38408", score=9.8, matches=[
            _match("cpe:2.3:a:openbsd:openssh:*:*:*:*:*:*:*:*", end_excl="9.3")]),
        _nvd_cve("CVE-2012-2122", score=5.1, severity="MEDIUM", matches=[
            _match("cpe:2.3:a:mysql:mysql:*:*:*:*:*:*:*:*", end_excl="5.1.63")]),
    )))
    cve_settings.put(db, cve_settings.WATERMARK, datetime(2026, 9, 30).isoformat())
    db.flush()


def test_findings_come_from_third_party_software_not_distro_packages(db):
    _load_cves(db)
    asset = _asset(db)
    store.save(db, asset.id, "linux", [
        _pkg("docker-ce", "5:24.0.7-1~ubuntu.22.04~jammy"),
        _pkg("openssh-server", "1:8.9p1-3ubuntu0.1", source="distro", origin="Ubuntu", source_package="openssh"),
    ])
    rows = [f for f in findings.compute(db, asset.id)["findings"] if f["asset_id"] == asset.id]
    assert {f["cve_id"] for f in rows} == {"CVE-2024-41110"}
    docker = rows[0]
    assert docker["identity_source"] == "inventory" and docker["installed"] == "24.0.7"


def test_firmware_and_service_versions_are_identified_and_matched(db):
    cve_store.apply_cves(db, feeds.parse_nvd_page(nvd_page(FORTI)))
    cve_settings.put(db, cve_settings.WATERMARK, datetime(2026, 9, 30).isoformat())
    asset = _asset(db, os_name="FortiOS", os_version="7.2.4")
    hooks.save_single(db, asset.id, "fortinet", collect.fortinet_item({"fortios_version": "7.2.5"}))
    rows = [f for f in findings.compute(db, asset.id)["findings"] if f["asset_id"] == asset.id]
    # one finding, from the version the audit read - not a second one from the asset's own fields
    assert [(f["cve_id"], f["installed"], f["identity_source"]) for f in rows] == [("CVE-2024-21762", "7.2.5",
                                                                                     "inventory")]
    for item in (collect.cisco_item("Cisco IOS XE Software, Version 17.09.04a"),
                 collect.mongodb_item("db version v6.0.12"),
                 collect.mssql_item("Microsoft SQL Server 2019 (RTM-CU22) - 15.0.4322.2 (X64)")):
        r = identify.resolve(item, {})
        assert r.status == "known" and r.nvd, item["name"]
    assert identify.resolve(collect.mssql_item("Microsoft SQL Server 2019 (RTM) - 15.0.2000.5 (X64)"),
                            {}).label == "Microsoft SQL Server 2019"


def test_full_inventory_drops_product_guesses():
    asset = Asset(asset_name="db-01", os_name="Ubuntu", os_version="22.04", model="MySQL 8.0.33")
    guessed = {(i.vendor, i.product): i.source for i in cve_cpe.identities_for(asset, [])}
    assert guessed[("oracle", "mysql")] == "inferred"
    # the package list says MySQL is not installed: the guess goes, the OS stays
    after = {(i.vendor, i.product) for i in cve_cpe.identities_for(asset, [], inventory=[], full_inventory=True)}
    assert ("oracle", "mysql") not in after and ("canonical", "ubuntu_linux") in after
    # the inventory's own version replaces a guess for the same product
    inv = [("oracle", "mysql", "8.0.36", "MySQL")]
    ids = [i for i in cve_cpe.identities_for(asset, [], inventory=inv) if i.product == "mysql"]
    assert [(i.version, i.source) for i in ids] == [("8.0.36", "inventory")]


def test_product_view_counts_versions_assets_and_findings(db):
    _load_cves(db)
    a1, a2 = _asset(db), _asset(db)
    store.save(db, a1.id, "linux", [_pkg("docker-ce", "5:24.0.7-1"), _pkg("docker-ce-cli", "5:24.0.7-1",
                                                                         source_package="docker-ce")])
    store.save(db, a2.id, "linux", [_pkg("docker-ce", "5:27.3.1-1"),
                                    _pkg("acme-monitor", "2.3.1", source="manual", origin=None)])
    data = service.products(db)
    docker = next(r for r in data["items"] if r["label"] == "Docker Engine")
    assert docker["asset_count"] == 2 and {v["version"] for v in docker["versions"]} == {"5:24.0.7-1", "5:27.3.1-1"}
    assert docker["cve"]["count"] == 1 and docker["cve"]["affected_versions"] == ["24.0.7"]
    acme = next(r for r in data["items"] if r["label"] == "acme-monitor")
    assert acme["status"] == "unknown"
    assert data["counts"]["unidentified"] >= 1 and data["counts"]["multi_version"] >= 1

    inv = service.asset_inventory(db, a1.id)
    assert inv["summary"]["vulnerable"] == 2               # both docker packages carry the product's findings
    overview = {r["asset_id"]: r for r in service.assets_overview(db)}
    assert overview[a1.id]["cves"] == 1 and overview[a2.id]["unidentified"] == 1


# ── API ───────────────────────────────────────────────────────────────────

def test_mapping_validation_and_effect(db):
    admin = _user(db)
    asset = _asset(db)
    store.save(db, asset.id, "linux", [_pkg("acme-monitor", "2.3.1", source="manual", origin=None)])
    with pytest.raises(HTTPException) as e:
        api.put_map(api.MapIn(match_key="acme-monitor", status="cpe", vendor="Acme Corp", product="x"), db, admin)
    assert e.value.status_code == 400
    out = api.put_map(api.MapIn(match_key="acme-monitor", status="cpe", vendor="acme", product="monitor"), db, admin)
    assert out["vendor"] == "acme"
    item = next(i for i in service.asset_inventory(db, asset.id)["items"] if i["name"] == "acme-monitor")
    assert item["status"] == "known" and item["cpes"] == ["acme:monitor"]
    api.delete_map(out["id"], db, admin)
    assert db.query(SoftwareProductMap).filter_by(match_key="acme-monitor").count() == 0


def test_collect_now_never_returns_401_and_logs(db, monkeypatch):
    from app.core.ssh_exceptions import SSHAuthenticationError
    from app.models.security_audit_log import AuditLog as SecurityAuditLog
    user = _user(db)
    asset = _asset(db)

    def refuse(*a, **kw):
        raise SSHAuthenticationError("Authentication failed", asset.ip_address)

    monkeypatch.setattr(hooks, "collect_now_linux", refuse)
    body = api.CollectIn(platform="linux", username="auditor", password="secret")
    with pytest.raises(HTTPException) as e:
        asyncio.run(api.collect_now(asset.id, body, db, user))
    assert e.value.status_code == 400
    log = db.query(SecurityAuditLog).filter_by(action="software.collect").order_by(SecurityAuditLog.id.desc()).first()
    assert log.result == "failed" and "secret" not in (log.detail or "")

    def ok(db_, asset_, **kw):
        assert kw["password"] == "secret" and kw["port"] == 22
        return store.save(db_, asset_.id, "linux", [_pkg("nginx", "1.26.0", origin="nginx.org")], trigger="manual")

    monkeypatch.setattr(hooks, "collect_now_linux", ok)
    monkeypatch.setattr(api, "_refresh_risk", lambda db_, aid: None)
    res = asyncio.run(api.collect_now(asset.id, body, db, user))
    assert res["collection"]["trigger"] == "manual" and res["collection"]["items"] == 1


def test_excel_export(db):
    asset = _asset(db)
    store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7")])
    from openpyxl import load_workbook
    from app.modules.software import export
    wb = load_workbook(export.products_workbook(service.products(db), "fa"))
    ws = wb.active
    assert ws.sheet_view.rightToLeft and ws["A1"].value == "محصول"
    assert any(row[0] == "Docker Engine" for row in ws.iter_rows(min_row=2, values_only=True))


# ── alert rule ────────────────────────────────────────────────────────────

def test_alert_new_software_outside_known_repositories(db):
    asset = _asset(db)
    since = datetime.utcnow() - timedelta(minutes=1)
    store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7")])
    # docker's repository was already there: a new package from it is not unexpected
    store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7"), _pkg("containerd.io", "1.6.26")])
    event = alert_events.EVENTS["software.unexpected"]
    assert event.enabled is False and event.module == "software"
    assert [p for p in event.evaluate(db, {}, datetime.utcnow(), since) if p.asset_id == asset.id] == []

    col = store.save(db, asset.id, "linux", [_pkg("docker-ce", "24.0.7"), _pkg("containerd.io", "1.6.26"),
                                             _pkg("acme-monitor", "2.3.1", source="manual", origin=None),
                                             _pkg("zabbix-agent2", "6.4.10", origin="repo.zabbix.com")])
    problems = [p for p in event.evaluate(db, {}, datetime.utcnow(), since) if p.asset_id == asset.id]
    assert len(problems) == 1 and problems[0].key == f"software:{col.id}"
    assert "acme-monitor" in problems[0].detail and "repo.zabbix.com" in problems[0].detail
    assert problems[0].link == f"/assets/software?asset={asset.id}"
    assert db.query(SoftwareCollection).filter_by(asset_id=asset.id).count() == 3
