"""Tests for Excel import (Asset List -> Import), app/utils/excel_utils.py.

The template/export this import must round-trip against spans 4 sheets
(Overview, Network & System, Location & Owner, Security & Audit), each keyed
by "Asset Name" - see ASSET_LIST_SHEETS in excel_utils.py. The import used to
only read wb.active (one sheet) with a full single-sheet header set that
didn't match any of the 4 sheets' actual (abbreviated) headers, so a file
downloaded from Export/Download Template and re-uploaded unmodified could
never successfully import.

Needs a migrated PostgreSQL database; each test runs inside a transaction
that is rolled back, so nothing here touches real rows (same pattern as
test_asset_deletion.py). import_assets_from_excel commits internally, which
lands on a savepoint here.
"""

import sys
from io import BytesIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pytest
from openpyxl import Workbook
from sqlalchemy import event
from sqlalchemy.orm import Session

from app.core.database import engine
from app.models.asset import Asset
from app.models.asset_types import AssetType
from app.models.user import User
from app.utils.excel_utils import ASSET_LIST_SHEETS, import_assets_from_excel


@pytest.fixture
def db():
    connection = engine.connect()
    trans = connection.begin()
    session = Session(bind=connection)
    session.begin_nested()

    @event.listens_for(session, "after_transaction_end")
    def _restart_savepoint(sess, transaction):
        if transaction.nested and not transaction._parent.nested:
            sess.begin_nested()

    try:
        yield session
    finally:
        event.remove(session, "after_transaction_end", _restart_savepoint)
        session.close()
        trans.rollback()
        connection.close()


_seq = [0]


def _next() -> int:
    _seq[0] += 1
    return _seq[0]


@pytest.fixture
def user(db) -> User:
    row = User(username=f"excel_import_user_{_next()}", hashed_password="x")
    db.add(row)
    db.flush()
    return row


@pytest.fixture
def cisco_router_type(db) -> AssetType:
    row = AssetType(type_name=f"Cisco Router {_next()}", category="network")
    db.add(row)
    db.flush()
    return row


def _multi_sheet_workbook(sheet_rows: dict) -> BytesIO:
    """Build a workbook shaped exactly like ASSET_LIST_SHEETS: one sheet per
    entry in `sheet_rows`, headers taken from ASSET_LIST_SHEETS, one data row
    appended per sheet title present in `sheet_rows`."""
    columns_by_title = {s["title"]: s["columns"] for s in ASSET_LIST_SHEETS}
    wb = Workbook()
    wb.remove(wb.active)
    for title, row in sheet_rows.items():
        ws = wb.create_sheet(title)
        ws.append(columns_by_title[title])
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestMultiSheetImport:
    def test_fields_from_all_four_sheets_are_merged_into_one_asset(self, db, user, cisco_router_type):
        wb = _multi_sheet_workbook({
            "Overview": ["Merged-Asset", "merged-host", cisco_router_type.type_name, "Core Router", "Cisco", "ISR4321"],
            "Network & System": ["Merged-Asset", "SN999", "IOS-XE", "10.0.0.77", "AA:BB:CC:DD:EE:FF", "24"],
            "Location & Owner": ["Merged-Asset", None, None, "active"],
            "Security & Audit": ["Merged-Asset", "internal", "low", None, None],
        })

        results = import_assets_from_excel(wb, db, user)

        assert results["errors"] == []
        assert results["created"] == 1
        asset = db.query(Asset).filter(Asset.asset_name == "Merged-Asset").first()
        assert asset is not None
        assert asset.hostname == "merged-host"
        assert asset.asset_type_id == cisco_router_type.id
        assert asset.asset_role == "Core Router"
        assert asset.manufacturer == "Cisco"
        assert asset.model == "ISR4321"
        assert asset.ip_address == "10.0.0.77"
        assert asset.mac_address == "AA:BB:CC:DD:EE:FF"

    def test_asset_type_from_overview_sheet_resolves_by_name(self, db, user, cisco_router_type):
        """The historical bug: Overview's "Type" header must map to the same
        field the "Asset Type" foreign-key lookup uses, or every new asset
        fails with "Asset type is required for new assets" even though a
        type was filled in."""
        wb = _multi_sheet_workbook({
            "Overview": ["Typed-Asset", None, cisco_router_type.type_name, None, None, None],
        })

        results = import_assets_from_excel(wb, db, user)

        assert results["errors"] == []
        assert results["created"] == 1
        asset = db.query(Asset).filter(Asset.asset_name == "Typed-Asset").first()
        assert asset.asset_type_id == cisco_router_type.id

    def test_unknown_asset_type_is_reported_not_silently_dropped(self, db, user):
        wb = _multi_sheet_workbook({
            "Overview": ["Bad-Type-Asset", None, "Nonexistent Type Xyz", None, None, None],
        })

        results = import_assets_from_excel(wb, db, user)

        assert results["created"] == 0
        assert any("Nonexistent Type Xyz" in e and "not found" in e for e in results["errors"])

    def test_new_asset_without_any_type_is_rejected(self, db, user):
        wb = _multi_sheet_workbook({
            "Network & System": ["No-Type-Asset", "SN1", None, "10.0.0.1", None, None],
        })

        results = import_assets_from_excel(wb, db, user)

        assert results["created"] == 0
        assert any("Asset type is required" in e for e in results["errors"])

    def test_reimporting_an_unmodified_export_updates_rather_than_errors(self, db, user, cisco_router_type):
        """The actual reported bug: download the template/export, fill it in,
        upload it back unmodified - it must succeed on both the first
        (create) and second (update) pass."""
        wb1 = _multi_sheet_workbook({
            "Overview": ["Round-Trip", "rt-host", cisco_router_type.type_name, None, None, None],
        })
        first = import_assets_from_excel(wb1, db, user)
        assert first["created"] == 1
        assert first["errors"] == []

        wb2 = _multi_sheet_workbook({
            "Overview": ["Round-Trip", "rt-host-renamed", cisco_router_type.type_name, None, None, None],
        })
        second = import_assets_from_excel(wb2, db, user)
        assert second["updated"] == 1
        assert second["created"] == 0
        assert second["errors"] == []

        asset = db.query(Asset).filter(Asset.asset_name == "Round-Trip").first()
        assert asset.hostname == "rt-host-renamed"


class TestLegacySingleSheetImport:
    """A hand-built or older-format file with one sheet and full header names
    (no sheet titled Overview/Network & System/etc present at all) must still
    work via the fallback."""

    def test_legacy_single_sheet_still_imports(self, db, user, cisco_router_type):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Asset Name", "Asset Type", "Hostname"])
        ws.append(["Legacy-Asset", cisco_router_type.type_name, "legacy-host"])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        results = import_assets_from_excel(buf, db, user)

        assert results["errors"] == []
        assert results["created"] == 1
        asset = db.query(Asset).filter(Asset.asset_name == "Legacy-Asset").first()
        assert asset is not None
        assert asset.asset_type_id == cisco_router_type.id
        assert asset.hostname == "legacy-host"
