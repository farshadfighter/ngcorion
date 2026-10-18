"""
Compliance with standards (app/modules/compliance): checks put in security domains, the built-in
standards (PCI DSS, ISO/IEC 27001, NIST SP 800-53) loaded and kept in step, requirements scored from the
latest audits (per asset and benchmark, without policy-excluded or informational results), assessed by
hand only where audits give no evidence, corrected check domains, gaps, the organisation's own standards
(import, export, edit), the audit trail and the guarded routes.

Needs a migrated PostgreSQL database; every test runs in a transaction that is rolled back, and works on
the assets of its own group so other audits in the database do not count.
"""
import json
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.core.database import engine
from app.models import Asset, AssetType, User, UserRole
from app.models.asset_groups import AssetGroup, AssetGroupMember
from app.models.audit import AuditResult, AuditSession, CheckStatus
from app.models.compliance import ComplianceAssessment, ComplianceRequirement, ComplianceStandard
from app.models.security_audit_log import AuditLog as SecurityAuditLog
from app.modules.compliance import domains as D
from app.modules.compliance import router as R
from app.modules.compliance import service as S

P_, F_ = CheckStatus.PASS, CheckStatus.FAIL


@pytest.fixture
def db():
    connection = engine.connect()
    trans = connection.begin()
    session = Session(bind=connection, join_transaction_mode="create_savepoint")
    S._synced = False
    try:
        yield session
    finally:
        session.close()
        trans.rollback()
        connection.close()
        S._synced = False


_seq = [0]


def _n():
    _seq[0] += 1
    return _seq[0]


def _user(db, role=UserRole.ADMIN):
    u = User(username=f"cmp_{_n()}", hashed_password="x", role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _group(db, *assets):
    g = AssetGroup(name=f"cmp group {_n()}", kind="static")
    db.add(g)
    db.flush()
    for a in assets:
        db.add(AssetGroupMember(group_id=g.id, asset_id=a.id))
    db.flush()
    return g


def _asset(db):
    t = AssetType(type_name=f"Server {_n()}", category="server")
    db.add(t)
    db.flush()
    a = Asset(asset_name=f"cmp-{_n()}", asset_type_id=t.id, ip_address=f"10.67.0.{_n() % 250}")
    db.add(a)
    db.flush()
    return a


LOCKOUT = "Ensure lockout for failed password attempts is configured"
COMPLEXITY = "Ensure password complexity is configured"
AUDITD = "Ensure auditd service is enabled"
CIPHERS = "Ensure sshd Ciphers are configured"
NTP = "Ensure chrony is configured"

# check, title, severity, level, status
RESULTS = [("5.3.1", LOCKOUT, "medium", "L1", P_), ("5.3.2", COMPLEXITY, "high", "L1", F_),
           ("6.2.1", AUDITD, "medium", "L1", P_), ("5.1.6", CIPHERS, "high", "L1", F_),
           ("2.3.1", NTP, "low", "L1", P_), ("9.9", "Ensure auditd logs are reviewed", "info", "INFO", F_)]


def _audit(db, asset, user, results=RESULTS, when=None, device_type="linux"):
    when = when or datetime.utcnow()
    s = AuditSession(user_id=user.id, asset_id=asset.id, target_ip=asset.ip_address, device_type=device_type,
                     status="completed", started_at=when, completed_at=when, compliance_pct=50.0)
    db.add(s)
    db.flush()
    for check, title, sev, level, st in results:
        db.add(AuditResult(session_id=s.id, check_number=check, check_title=title, severity=sev, level=level,
                           status=st))
    db.flush()
    return s


def _req():
    return SimpleNamespace(client=SimpleNamespace(host="10.0.0.9"))


def _std(db, key="pci-dss-4.0"):
    S.sync_builtin(db)
    return db.query(ComplianceStandard).filter_by(key=key).one()


def _by_ref(db, std, group, user):
    out = R.get_standard(std.id, group_id=group.id, current_user=user, db=db)
    return {r["ref"]: r for r in out["requirements"]}, out["summary"]


def _setup(db):
    user = _user(db)
    a = _asset(db)
    _audit(db, a, user)
    return user, a, _group(db, a)


# ── domains ──────────────────────────────────────────────────────────────

def test_titles_are_put_in_domains_and_topics():
    assert "authentication" in D.classify(LOCKOUT)
    assert "lockout" in D.topics(LOCKOUT)
    assert D.classify(AUDITD) == ["logging"]
    assert set(D.classify(CIPHERS)) == {"remote_access", "crypto"}
    assert D.classify(NTP) == ["time"]
    assert D.classify("Ensure net.ipv4.conf.default.forwarding is disabled") == ["network_hardening"]
    # "log on" is not logging; a Windows user right is an account matter by its section
    assert D.classify("Ensure 'Shutdown: Allow system to be shut down without having to log on'") == []
    assert D.classify("Ensure 'Back up files and directories'", "windows", "WIN-2025-2.2.7") == ["accounts"]


def test_a_requirement_names_domains_both_domains_topics_or_checks():
    m = {"domains": ["remote_access+crypto"]}
    assert D.requirement_matches(m, "linux", "5.1.6", ["remote_access", "crypto"], [])
    assert not D.requirement_matches(m, "linux", "x", ["crypto"], [])
    assert D.requirement_matches({"topics": ["lockout"]}, "linux", "x", [], ["lockout"])
    assert D.requirement_matches({"checks": ["linux:5.3.2"]}, "linux", "5.3.2", [], [])
    assert not D.requirement_matches({"checks": ["cisco:5.3.2"]}, "linux", "5.3.2", [], [])
    assert D.requirement_matches({"checks": ["5.3.2"]}, "cisco", "5.3.2", [], [])
    assert D.requirement_matches({"all": True}, "linux", "x", [], [])
    assert not D.requirement_matches({}, "linux", "x", ["logging"], [])
    assert D.validate_mapping({"domains": ["logging+nope"]}) == "Unknown domain: nope"
    assert D.validate_mapping({"topics": ["nope"]}) == "Unknown topic: nope"


def test_builtin_standards_are_valid_and_loaded_once(db):
    for data in S.builtin_files():
        refs = [r["ref"] for r in data["requirements"]]
        assert len(refs) == len(set(refs)), data["key"]
        sections = {s["ref"] for s in data["sections"]}
        for r in data["requirements"]:
            assert r["section"] in sections and r["title_fa"], (data["key"], r["ref"])
            assert D.validate_mapping(r["mapping"]) is None, (data["key"], r["ref"])
    S.sync_builtin(db)
    stds = {s.key: s for s in db.query(ComplianceStandard).filter_by(builtin=True)}
    assert {"pci-dss-4.0", "iso-27001-2022", "nist-800-53-r5"} <= set(stds)
    n = db.query(ComplianceRequirement).filter_by(standard_id=stds["pci-dss-4.0"].id).count()
    assert n == len(next(d for d in S.builtin_files() if d["key"] == "pci-dss-4.0")["requirements"])
    # the organisation's choice is kept when the definitions are applied again
    stds["iso-27001-2022"].enabled = False
    db.flush()
    S._synced = False
    S.sync_builtin(db)
    db.refresh(stds["iso-27001-2022"])
    assert stds["iso-27001-2022"].enabled is False
    assert db.query(ComplianceRequirement).filter_by(standard_id=stds["pci-dss-4.0"].id).count() == n


def test_every_check_of_every_module_has_a_domain_mostly():
    """Most checks of the 17 modules land in a domain (the rest count for secure configuration only)."""
    from app.modules.benchmark import registry
    titles = [(s.device_type.value, r.id, r.title) for s in registry.specs() for r in s.all_rules()]
    from app.modules.linux.audit.rules import get_linux_rule_catalog
    titles += [("linux", r.id if hasattr(r, "id") else "", r.title) for r in get_linux_rule_catalog().values()]
    classified = sum(1 for b, n, t in titles if D.classify(t, b, n))
    assert len(titles) > 1000
    assert classified / len(titles) > 0.9


# ── scoring ──────────────────────────────────────────────────────────────

def test_requirements_scored_from_the_latest_audits(db):
    user, a, g = _setup(db)
    std = _std(db)
    reqs, summary = _by_ref(db, std, g, user)
    assert reqs["8.3.4"]["status"] == "met" and reqs["8.3.4"]["score"] == 100.0       # lockout passed
    assert reqs["8.3.6"]["status"] == "not_met" and reqs["8.3.6"]["score"] == 0.0     # complexity failed
    assert reqs["10.2.1"]["status"] == "met"                                           # auditd
    assert reqs["10.6.1"]["status"] == "met"                                           # chrony
    assert reqs["2.2.7"]["status"] == "not_met"                                        # ssh ciphers
    assert reqs["2.2.1"]["status"] == "partial" and reqs["2.2.1"]["score"] == 60.0     # every check: 3 of 5
    assert reqs["12.1.1"]["status"] == "not_assessed" and reqs["12.1.1"]["manual_only"]
    assert summary["counts"]["not_assessed"] > 0 and 0 < summary["coverage"] < 100
    # the informational result is not counted
    assert reqs["10.2.1"]["passed"] + reqs["10.2.1"]["failed"] == 1


def test_only_the_latest_audit_of_each_benchmark_counts_and_policy_results_do_not(db):
    user, a = _user(db), _asset(db)
    _audit(db, a, user, results=[("5.3.2", COMPLEXITY, "high", "L1", P_)], when=datetime.utcnow() - timedelta(days=30))
    _audit(db, a, user)
    g = _group(db, a)
    # the same asset as a Docker host counts too
    _audit(db, a, user, device_type="docker",
           results=[("DKR-1", "Ensure password complexity is configured for the registry", "high", "L1", P_)])
    std = _std(db)
    reqs, _ = _by_ref(db, std, g, user)
    assert (reqs["8.3.6"]["passed"], reqs["8.3.6"]["failed"]) == (1, 1)
    latest = db.query(AuditSession).filter_by(asset_id=a.id, device_type="linux").order_by(AuditSession.id.desc()).first()
    for r in db.query(AuditResult).filter_by(session_id=latest.id, check_number="5.3.2"):
        r.policy = "excluded"
    db.flush()
    reqs, _ = _by_ref(db, std, g, user)
    assert (reqs["8.3.6"]["passed"], reqs["8.3.6"]["failed"]) == (1, 0)


def test_assessed_by_hand_only_without_evidence(db):
    user, a, g = _setup(db)
    std = _std(db)
    reqs, _ = _by_ref(db, std, g, user)
    policy_req, complexity_req = reqs["12.1.1"]["id"], reqs["8.3.6"]["id"]
    R.assess(policy_req, R.AssessmentIn(group_id=g.id, status="met", note="Policy approved by the board"),
             request=_req(), current_user=user, db=db)
    with pytest.raises(HTTPException) as e:
        R.assess(complexity_req, R.AssessmentIn(group_id=g.id, status="met", note="We say so"),
                 request=_req(), current_user=user, db=db)
    assert e.value.status_code == 400
    R.assess(complexity_req, R.AssessmentIn(group_id=g.id, status="not_applicable", note="No local accounts here"),
             request=_req(), current_user=user, db=db)
    reqs, _ = _by_ref(db, std, g, user)
    assert reqs["12.1.1"]["status"] == "met" and reqs["12.1.1"]["source"] == "manual"
    assert reqs["12.1.1"]["assessment"]["assessed_by"] == user.username
    assert reqs["8.3.6"]["status"] == "not_applicable" and reqs["8.3.6"]["score"] is None
    # an assessment for all assets does not count for the group
    other = R.get_standard(std.id, group_id=None, current_user=user, db=db)
    assert next(r for r in other["requirements"] if r["ref"] == "12.1.1")["status"] == "not_assessed"
    # expired
    a_ = db.query(ComplianceAssessment).filter_by(requirement_id=policy_req).one()
    a_.valid_until = datetime.utcnow() - timedelta(days=1)
    db.flush()
    reqs, _ = _by_ref(db, std, g, user)
    assert reqs["12.1.1"]["status"] == "not_assessed" and reqs["12.1.1"]["assessment_expired"]
    R.remove_assessment(policy_req, request=_req(), group_id=g.id, current_user=user, db=db)
    assert not db.query(ComplianceAssessment).filter_by(requirement_id=policy_req).count()


def test_an_administrator_corrects_the_domains_of_a_check(db):
    user, a, g = _setup(db)
    std = _std(db)
    before, _ = _by_ref(db, std, g, user)
    assert before["10.6.1"]["checks"] == 1
    manager = _user(db, UserRole.MANAGER)
    body = R.DomainsIn(benchmark="linux", check_number="2.3.1", domains=["logging"])
    with pytest.raises(HTTPException) as e:
        R.set_check_domains(body, request=_req(), current_user=manager, db=db)
    assert e.value.status_code == 403
    R.set_check_domains(body, request=_req(), current_user=user, db=db)
    after, _ = _by_ref(db, std, g, user)
    assert after["10.6.1"]["checks"] == 0 and after["10.2.1"]["checks"] == 2
    row = next(c for c in R.list_checks(benchmark=None, q="2.3.1", domain=None, unclassified=False, offset=0, limit=100,
                                       current_user=user, db=db)["checks"] if c["benchmark"] == "linux")
    assert row["overridden"] and row["domains"] == ["logging"]
    R.set_check_domains(R.DomainsIn(benchmark="linux", check_number="2.3.1", domains=None), request=_req(),
                        current_user=user, db=db)
    again, _ = _by_ref(db, std, g, user)
    assert again["10.6.1"]["checks"] == 1


def test_gaps_and_failing_assets(db):
    user, a, g = _setup(db)
    std = _std(db)
    gaps = R.list_gaps(std.id, group_id=g.id, limit=200, current_user=user, db=db)["gaps"]
    by = {x["check_number"]: x for x in gaps}
    assert set(by) == {"5.3.2", "5.1.6"}
    assert "8.3.6" in by["5.3.2"]["requirements"] and by["5.3.2"]["requirements"][-1] == "2.2.1"
    assets = R.list_failing_assets(benchmark="linux", check_number="5.3.2", group_id=g.id, current_user=user,
                                   db=db)["assets"]
    assert [x["asset_id"] for x in assets] == [a.id]
    detail = R.get_requirement(by_ref_id(db, std, "8.3.6"), group_id=g.id, current_user=user, db=db)
    assert detail["evidence"][0]["check_number"] == "5.3.2" and detail["evidence"][0]["failed"] == 1


def by_ref_id(db, std, ref):
    return db.query(ComplianceRequirement).filter_by(standard_id=std.id, ref=ref).one().id


# ── the organisation's own standards ─────────────────────────────────────

def _definition():
    return {"name": "National baseline", "version": "1403", "sections": [{"ref": "A", "title": "Access"}],
            "requirements": [
                {"ref": "A-1", "section": "A", "title": "Lock accounts after failed logons",
                 "title_fa": "قفل حساب پس از ورود ناموفق", "mapping": {"topics": ["lockout"]}},
                {"ref": "A-2", "section": "A", "title": "Remote administration over encrypted channels",
                 "mapping": {"domains": ["remote_access+crypto"]}},
                {"ref": "A-3", "section": "A", "title": "This one check", "mapping": {"checks": ["linux:5.3.2"]}},
                {"ref": "B-1", "title": "Security committee meets every quarter"}]}


def test_import_score_export_and_edit_an_own_standard(db):
    user, a, g = _setup(db)
    std = R.import_standard(R.StandardImport(**_definition()), request=_req(), current_user=user, db=db)
    reqs, summary = _by_ref(db, db.get(ComplianceStandard, std["id"]), g, user)
    assert reqs["A-1"]["status"] == "met" and reqs["A-2"]["status"] == "not_met"
    assert reqs["A-3"]["checks"] == 1 and reqs["B-1"]["manual_only"]
    assert summary["requirements"] == 4
    exported = R.export_standard(std["id"], current_user=user, db=db)
    assert [r["ref"] for r in exported["requirements"]] == ["A-1", "A-2", "A-3", "B-1"]
    json.dumps(exported)
    rid = R.add_requirement(std["id"], R.RequirementIn(ref="A-4", section="A", title="Time is synchronised",
                                                       mapping=R.MappingIn(domains=["time"])),
                            request=_req(), current_user=user, db=db)["id"]
    with pytest.raises(HTTPException) as e:
        R.add_requirement(std["id"], R.RequirementIn(ref="A-4", title="Again"), request=_req(), current_user=user,
                          db=db)
    assert e.value.status_code == 409
    R.update_requirement(rid, R.RequirementIn(ref="A-4", section="A", title="Clocks", mapping=R.MappingIn(
        domains=["logging"])), request=_req(), current_user=user, db=db)
    assert db.get(ComplianceRequirement, rid).mapping == {"domains": ["logging"]}
    R.delete_requirement(rid, request=_req(), current_user=user, db=db)
    R.update_standard(std["id"], R.StandardPatch(name="National baseline 1403"), request=_req(), current_user=user,
                      db=db)
    R.delete_standard(std["id"], request=_req(), current_user=user, db=db)
    assert db.get(ComplianceStandard, std["id"]) is None


def test_bad_definitions_and_builtin_standards_are_refused(db):
    user = _user(db)
    bad = _definition()
    bad["requirements"][0]["mapping"] = {"domains": ["telepathy"]}
    with pytest.raises(HTTPException) as e:
        R.import_standard(R.StandardImport(**bad), request=_req(), current_user=user, db=db)
    assert e.value.status_code == 400
    dup = _definition()
    dup["requirements"][1]["ref"] = "A-1"
    with pytest.raises(HTTPException):
        R.import_standard(R.StandardImport(**dup), request=_req(), current_user=user, db=db)
    pci = _std(db)
    with pytest.raises(HTTPException) as e:
        R.update_standard(pci.id, R.StandardPatch(name="Mine"), request=_req(), current_user=user, db=db)
    assert e.value.status_code == 400
    with pytest.raises(HTTPException):
        R.delete_standard(pci.id, request=_req(), current_user=user, db=db)
    R.update_standard(pci.id, R.StandardPatch(enabled=False), request=_req(), current_user=user, db=db)
    assert pci.id not in {s["id"] for s in R.list_standards(group_id=None, include_disabled=False, current_user=user, db=db)["standards"]}
    viewer = _user(db, UserRole.USER)
    with pytest.raises(HTTPException) as e:
        R.create_standard(R.StandardIn(name="X1"), request=_req(), current_user=viewer, db=db)
    assert e.value.status_code == 403


def test_the_compliance_report_builds_with_the_group_and_standard(db):
    import io

    from openpyxl import load_workbook

    from app.modules.reports import service as reports
    user, a, g = _setup(db)
    params = {"period": {"preset": "custom", "from": "2026-09-01", "to": "2026-09-30"}, "compare": False,
              "scope": {"mode": "all"}, "options": {"standard": "nist-800-53-r5", "group": str(g.id)},
              "sections": ["overview", "requirements", "gaps", "manual", "failed_all"]}
    r = reports.create(db, user, "compliance_standard", "", "fa", ["pdf", "xlsx"], "internal", params)
    out = reports.build(db, r, now=datetime.utcnow())
    assert out["files"]["pdf"][1].startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(out["files"]["xlsx"][1]))
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
    assert "AC-7" in text and "IA-5" in text           # requirements
    assert "5.3.2" in text and a.asset_name in text    # gaps and every failed check
    assert "تلاش‌های ناموفق ورود" in text


def test_changes_are_in_the_audit_trail(db):
    user, a, g = _setup(db)
    std = R.import_standard(R.StandardImport(**_definition()), request=_req(), current_user=user, db=db)
    rows = db.query(SecurityAuditLog).filter(SecurityAuditLog.user_id == user.id).all()
    assert any(r.action == "compliance.standard_import" and "National baseline" in r.detail
               and "4 requirements" in r.detail for r in rows)
    assert std


def test_routes_are_guarded():
    assert len(R.router.routes) >= 15
    for route in R.router.routes:
        deps = [d.call for d in route.dependant.dependencies]
        assert any(getattr(c, "permission", None) for c in deps), route.path
