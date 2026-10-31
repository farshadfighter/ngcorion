"""
TLS certificates (app/modules/tls): reading real handshakes from local TLS servers (trust through the
public list or the organisation's own authority, self-signed, expired, weak key and SHA-1, TLS 1.0 still
accepted, no answer, not TLS), judging a certificate (expiry, trust, name, key, signature, protocol),
choosing what is scanned and keeping the services, the API (who may change what), and what follows:
remediation items, the risk score, alerts, the audit trail and the report.

Needs a migrated PostgreSQL database; every test runs in a transaction that is rolled back.
"""
import socket
import ssl
import tempfile
import threading
import warnings
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from cryptography import x509
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import ec, rsa
from cryptography.x509.oid import NameOID
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.core.database import engine
from app.models import Asset, AssetType, User, UserRole
from app.models.enums import StatusEnum
from app.models.risk import AssetOpenPort
from app.models.security_audit_log import AuditLog as SecurityAuditLog
from app.models.tls import SOURCE_AUTO, SOURCE_MANUAL, TlsEndpoint, TlsTrustedCa
from app.modules.tls import checks as K
from app.modules.tls import router as R
from app.modules.tls import scanner
from app.modules.tls import service as S

NOW = datetime(2026, 10, 18, 12, 0)


@pytest.fixture
def db():
    connection = engine.connect()
    trans = connection.begin()
    session = Session(bind=connection, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        session.close()
        trans.rollback()
        connection.close()


_seq = [0]


def _n():
    _seq[0] += 1
    return _seq[0]


def _user(db, role=UserRole.ADMIN):
    u = User(username=f"tls_{_n()}", hashed_password="x", role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _asset(db, ip=None, hostname=None, status=StatusEnum.ACTIVE):
    t = AssetType(type_name=f"Server {_n()}", category="server")
    db.add(t)
    db.flush()
    a = Asset(asset_name=f"tls-{_n()}", asset_type_id=t.id, ip_address=ip or f"10.69.0.{_n() % 250}",
              hostname=hostname, status=status)
    db.add(a)
    db.flush()
    return a


def _req():
    return SimpleNamespace(client=SimpleNamespace(host="10.0.0.9"))


# ── certificates and local TLS servers ───────────────────────────────────

def _name(cn):
    return x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, cn)])


def _ca(cn="Corp Root CA"):
    key = ec.generate_private_key(ec.SECP256R1())
    now = datetime.now(timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(_name(cn)).issuer_name(_name(cn)).public_key(key.public_key())
            .serial_number(x509.random_serial_number()).not_valid_before(now - timedelta(days=30))
            .not_valid_after(now + timedelta(days=3650))
            .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
            .add_extension(x509.KeyUsage(False, False, False, False, False, True, True, False, False), critical=True)
            .sign(key, hashes.SHA256()))
    return cert, key


def _leaf(cn="srv.corp.local", issuer=None, days=(-10, 365), rsa_bits=None, sans=None):
    key = rsa.generate_private_key(65537, rsa_bits) if rsa_bits else ec.generate_private_key(ec.SECP256R1())
    now = NOW.replace(tzinfo=timezone.utc)          # the tests evaluate at NOW, not today
    ca_cert, ca_key = issuer if issuer else (None, key)
    b = (x509.CertificateBuilder().subject_name(_name(cn))
         .issuer_name(ca_cert.subject if ca_cert else _name(cn)).public_key(key.public_key())
         .serial_number(x509.random_serial_number()).not_valid_before(now + timedelta(days=days[0]))
         .not_valid_after(now + timedelta(days=days[1]))
         .add_extension(x509.SubjectAlternativeName([x509.DNSName(n) for n in (sans or [cn])]), critical=False)
         .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True))
    cert = b.sign(ca_key, hashes.SHA256())
    return cert, key


def _pem(cert):
    return cert.public_bytes(serialization.Encoding.PEM).decode()


class Server:
    """A TLS server on 127.0.0.1 that only completes handshakes."""

    def __init__(self, cert, key, chain=(), min_version=None):
        self.ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        self.ctx.set_ciphers("ALL:@SECLEVEL=0")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", DeprecationWarning)
            self.ctx.minimum_version = min_version or ssl.TLSVersion.TLSv1_2
        with tempfile.NamedTemporaryFile("w", suffix=".pem", delete=False) as f:
            f.write(_pem(cert) + "".join(_pem(c) for c in chain))
            f.write(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                      serialization.NoEncryption()).decode())
            path = f.name
        self.ctx.load_cert_chain(path)
        self.sock = socket.create_server(("127.0.0.1", 0))
        self.port = self.sock.getsockname()[1]
        threading.Thread(target=self._serve, daemon=True).start()

    def _serve(self):
        while True:
            try:
                conn, _ = self.sock.accept()
            except OSError:
                return
            try:
                with self.ctx.wrap_socket(conn, server_side=True) as t:
                    t.recv(1)
            except (OSError, ssl.SSLError):
                pass
            finally:
                conn.close()

    def close(self):
        self.sock.close()


@pytest.fixture
def servers():
    made = []

    def make(*args, **kwargs):
        s = Server(*args, **kwargs)
        made.append(s)
        return s
    yield make
    for s in made:
        s.close()


# ── reading a handshake ──────────────────────────────────────────────────

def test_probe_trust_through_the_organisations_authority(servers):
    ca = _ca()
    cert, key = _leaf(issuer=ca)
    srv = servers(cert, key)
    p = scanner.probe("127.0.0.1", srv.port, check_protocols=False)
    assert p.trusted is False and not p.self_signed and p.trust_error
    p = scanner.probe("127.0.0.1", srv.port, cadata=_pem(ca[0]), check_protocols=True)
    assert p.trusted is True and p.trust_error is None and p.old_protocols == []
    d = scanner.describe(p.chain[0])
    assert d["names"] == ["srv.corp.local"] and d["key_type"] == "EC" and d["key_bits"] == 256
    assert d["signature"] == "sha256" and not d["self_issued"] and len(d["fingerprint"]) == 64


def test_probe_self_signed_expired_and_weak(servers):
    # (cryptography no longer signs with SHA-1: that check is in test_problems)
    cert, key = _leaf(cn="old.corp.local", days=(-400, -5), rsa_bits=1024)
    srv = servers(cert, key, min_version=ssl.TLSVersion.TLSv1)
    p = scanner.probe("127.0.0.1", srv.port)
    assert p.trusted is False and p.self_signed
    assert "TLSv1" in p.old_protocols and "TLSv1.1" in p.old_protocols
    d = scanner.describe(p.chain[0])
    assert (d["key_type"], d["key_bits"], d["self_issued"]) == ("RSA", 1024, True)


def test_an_expired_certificate_of_a_trusted_authority_is_still_trusted(servers):
    ca = _ca()
    cert, key = _leaf(issuer=ca, days=(-400, -1))
    srv = servers(cert, key)
    p = scanner.probe("127.0.0.1", srv.port, cadata=_pem(ca[0]), check_protocols=False)
    assert p.trusted is True       # the chain is good: only the date failed


def test_probe_no_answer_and_not_tls():
    closed = socket.create_server(("127.0.0.1", 0))
    port = closed.getsockname()[1]
    closed.close()
    with pytest.raises(scanner.ProbeError) as e:
        scanner.probe("127.0.0.1", port, timeout=2)
    assert e.value.kind == "unreachable"

    plain = socket.create_server(("127.0.0.1", 0))

    def answer():
        conn, _ = plain.accept()
        conn.sendall(b"SSH-2.0-OpenSSH_9.6\r\n" * 4)
        conn.close()
    threading.Thread(target=answer, daemon=True).start()
    with pytest.raises(scanner.ProbeError) as e:
        scanner.probe("127.0.0.1", plain.getsockname()[1], timeout=2, check_protocols=False)
    plain.close()
    assert e.value.kind == "no_tls"


def test_parse_ca():
    ca = _ca()
    out = scanner.parse_ca(_pem(ca[0]))
    assert out["subject"] == "CN=Corp Root CA" and out["pem"].startswith("-----BEGIN CERTIFICATE-----")
    with pytest.raises(ValueError):
        scanner.parse_ca("not a certificate" * 5)
    leaf, _ = _leaf(issuer=ca)
    with pytest.raises(ValueError, match="not a certificate authority"):
        scanner.parse_ca(_pem(leaf))


# ── judging ──────────────────────────────────────────────────────────────

@pytest.mark.parametrize("names, host, wanted, expected", [
    (["10.0.0.5"], "10.0.0.5", [], True),
    (["fw.corp.local"], "10.0.0.5", [], None),
    (["fw.corp.local"], "10.0.0.5", ["fw"], True),
    (["fw.corp.local"], "10.0.0.5", ["fw.corp.local"], True),
    (["*.corp.local"], "10.0.0.5", ["fw.corp.local"], True),
    (["*.corp.local"], "10.0.0.5", ["a.b.corp.local"], False),
    (["fortigate"], "10.0.0.5", ["fw-edge"], False),
    (["fw.corp.local"], "10.0.0.5", ["10.0.0.6"], None),
])
def test_name_ok(names, host, wanted, expected):
    assert K.name_ok(names, host, wanted) is expected


def _ep(**kw):
    base = dict(fingerprint="ab" * 32, host="10.0.0.5", server_name=None, names=["srv.corp.local"],
                not_before=NOW - timedelta(days=30), not_after=NOW + timedelta(days=300), trusted=True,
                self_signed=False, trust_error=None, key_type="EC", key_bits=256, signature="sha256",
                old_protocols=[])
    base.update(kw)
    return SimpleNamespace(**base)


def test_problems():
    assert K.problems(_ep(), NOW, 30, ["srv"]) == [] and K.level(_ep(), []) == "ok"
    soon = K.problems(_ep(not_after=NOW + timedelta(days=12)), NOW, 30)
    assert [p["code"] for p in soon] == ["expiring"] and soon[0]["days"] == 12
    bad = K.problems(_ep(not_after=NOW - timedelta(days=1), trusted=False, self_signed=True, key_type="RSA",
                         key_bits=1024, signature="sha1", old_protocols=["TLSv1", "TLSv1.1"]), NOW, 30, ["web"])
    codes = {p["code"]: p for p in bad}
    assert set(codes) == {"expired", "untrusted", "weak_key", "weak_signature", "old_protocol", "name_mismatch"}
    assert codes["expired"]["severity"] == "high" and codes["name_mismatch"]["severity"] == "low"
    assert codes["weak_key"]["key"] == "RSA 1024" and codes["weak_signature"]["algorithm"] == "SHA-1"
    assert codes["old_protocol"]["protocols"] == ["TLS 1.0", "TLS 1.1"]
    assert K.level(_ep(), bad) == "critical" and K.level(_ep(), soon) == "warning"
    assert K.problems(_ep(fingerprint=None), NOW, 30) == []


# ── what is scanned, and what is kept ────────────────────────────────────

def _fake(answers):
    """A probe answering from a table {(host, port): Probe | kind of failure}; anything else: no answer."""
    asked = []

    def probe(host, port, server_name=None, cadata=None, check_protocols=True, timeout=None):
        asked.append((host, port, server_name, bool(cadata)))
        a = answers.get((host, port), "unreachable")
        if isinstance(a, str):
            raise scanner.ProbeError(a, a)
        return a
    probe.asked = asked
    return probe


def _answer(trusted=True, **cert):
    ca = cert.pop("issuer", None)
    c, _ = _leaf(issuer=ca, **cert)
    return scanner.Probe(protocol="TLSv1.3", chain=[c.public_bytes(serialization.Encoding.DER)], trusted=trusted,
                         old_protocols=[])


def test_scan_finds_services_on_usual_and_open_ports(db):
    S.save_settings(db, {"ports": [443, 8443]}, None)
    a = _asset(db, ip="10.70.0.1", hostname="web.corp.local")
    gone = _asset(db, ip="10.70.0.2", status=StatusEnum.DECOMMISSIONED)
    db.add(AssetOpenPort(asset_id=a.id, ip_address=a.ip_address, port=4443, protocol="tcp",
                         service_name="ssl/https-alt", severity_score=10))
    db.add(AssetOpenPort(asset_id=a.id, ip_address=a.ip_address, port=22, protocol="tcp", service_name="ssh",
                         severity_score=10))
    db.flush()
    probe = _fake({("10.70.0.1", 443): _answer(cn="web.corp.local"), ("10.70.0.1", 4443): _answer(cn="x")})
    stats = S.run_scan(db, asset_ids=[a.id, gone.id], now=NOW, probe=probe)
    assert stats == {"probed": 3, "read": 2, "new": 2, "failed": 0, "dropped": 0}
    assert {(h, p) for h, p, _, _ in probe.asked} == {("10.70.0.1", 443), ("10.70.0.1", 8443), ("10.70.0.1", 4443)}
    assert all(sni == "web.corp.local" for _, _, sni, _ in probe.asked)
    eps = {e.port: e for e in db.query(TlsEndpoint).filter_by(asset_id=a.id)}
    assert set(eps) == {443, 4443} and eps[443].status == "read" and eps[443].source == SOURCE_AUTO
    rows = {r["endpoint"].port: r for r in S.evaluate(db, [a.id], now=NOW)}
    assert rows[443]["level"] == "ok" and [p["code"] for p in rows[4443]["problems"]] == ["name_mismatch"]


def test_scan_keeps_failures_drops_old_ones_and_skips_turned_off(db):
    S.save_settings(db, {"ports": [443]}, None)
    a = _asset(db, ip="10.70.1.1")
    S.run_scan(db, asset_ids=[a.id], now=NOW, probe=_fake({("10.70.1.1", 443): _answer()}))
    ep = db.query(TlsEndpoint).filter_by(asset_id=a.id).one()
    first = ep.fingerprint
    S.run_scan(db, asset_ids=[a.id], now=NOW + timedelta(days=1), probe=_fake({}))
    db.refresh(ep)
    assert ep.status == "unreachable" and ep.fingerprint == first      # the last certificate stays
    S.run_scan(db, asset_ids=[a.id], now=NOW + timedelta(days=2), probe=_fake({("10.70.1.1", 443): _answer()}))
    db.refresh(ep)
    assert ep.fingerprint != first and ep.cert_changed_at == NOW + timedelta(days=2)
    ep.enabled = False
    db.flush()
    probe = _fake({})
    S.run_scan(db, asset_ids=[a.id], now=NOW + timedelta(days=3), probe=probe)
    assert probe.asked == [] and S.evaluate(db, [a.id]) == []
    ep.enabled = True
    db.flush()
    stats = S.run_scan(db, asset_ids=[a.id], now=NOW + timedelta(days=40), probe=_fake({}))
    assert stats["dropped"] == 1 and db.query(TlsEndpoint).filter_by(asset_id=a.id).count() == 0


def test_manual_service_uses_its_own_name_and_the_trusted_authorities(db):
    S.save_settings(db, {"ports": [443]}, None)
    ca = _ca()
    db.add(TlsTrustedCa(name="Corp", **scanner.parse_ca(_pem(ca[0]))))
    a = _asset(db, ip="10.70.2.1")
    ep = TlsEndpoint(asset_id=a.id, host="portal.corp.local", port=9000, server_name="portal.corp.local",
                     source=SOURCE_MANUAL, first_seen_at=NOW)
    db.add(ep)
    db.flush()
    probe = _fake({("portal.corp.local", 9000): _answer(cn="portal.corp.local", issuer=ca)})
    S.run_scan(db, endpoint_ids=[ep.id], now=NOW, probe=probe)
    assert probe.asked == [("portal.corp.local", 9000, "portal.corp.local", True)]
    db.refresh(ep)
    assert ep.status == "read" and ep.issuer == "CN=Corp Root CA"


def test_due(db):
    S.save_settings(db, {"interval_hours": 24, "auto_scan": True}, None)
    S.set_state(db, last_run=NOW.isoformat(), requested=False)
    assert not S.due(db, NOW + timedelta(hours=23)) and S.due(db, NOW + timedelta(hours=24))
    S.save_settings(db, {"auto_scan": False}, None)
    assert not S.due(db, NOW + timedelta(days=5))
    S.set_state(db, requested=True)
    assert S.due(db, NOW)


def test_settings_are_clamped(db):
    out = S.save_settings(db, {"warn_days": 9999, "ports": [443, 443, "x", 70000, 8443]}, None)
    assert out["warn_days"] == 180 and out["ports"] == [443, 8443]


# ── the API ──────────────────────────────────────────────────────────────

def test_routes_are_guarded():
    assert len(R.router.routes) >= 12
    for route in R.router.routes:
        deps = [d.call for d in route.dependant.dependencies]
        assert any(getattr(c, "permission", None) for c in deps), route.path


def test_api_services_settings_and_authorities(db, monkeypatch):
    monkeypatch.setattr(scanner, "probe", _fake({("10.70.3.1", 9443): _answer(cn="app.corp.local", trusted=False)}))
    admin, manager = _user(db), _user(db, UserRole.MANAGER)
    a = _asset(db, ip="10.70.3.1", hostname="app")
    out = R.add_service(R.ServiceIn(asset_id=a.id, port=9443), request=_req(), current_user=manager, db=db)
    assert out["status"] == "read" and out["source"] == "manual" and out["level"] == "warning"
    assert [p["code"] for p in out["problems"]] == ["untrusted"]
    with pytest.raises(HTTPException) as e:
        R.add_service(R.ServiceIn(asset_id=a.id, port=9443), request=_req(), current_user=manager, db=db)
    assert e.value.status_code == 409
    with pytest.raises(ValueError):
        R.ServiceIn(asset_id=a.id, port=443, host="bad host!")

    over = R.overview(current_user=manager, db=db)
    mine = [s for s in over["services"] if s["asset_id"] == a.id]
    assert mine[0]["where"] == "10.70.3.1:9443" and over["summary"]["problems"]["untrusted"] >= 1
    off = R.change_service(out["id"], R.ServicePatch(enabled=False), request=_req(), current_user=manager, db=db)
    assert off["enabled"] is False and off["problems"] == []

    with pytest.raises(HTTPException) as e:
        R.put_settings(R.SettingsIn(warn_days=14, auto_scan=True, interval_hours=12, check_protocols=True,
                                    ports=[443]), request=_req(), current_user=manager, db=db)
    assert e.value.status_code == 403
    saved = R.put_settings(R.SettingsIn(warn_days=14, auto_scan=True, interval_hours=12, check_protocols=False,
                                        ports=[443, 636]), request=_req(), current_user=admin, db=db)
    assert saved["ports"] == [443, 636] and saved["check_protocols"] is False

    ca = _ca("Branch CA")
    with pytest.raises(HTTPException):
        R.add_ca(R.CaIn(name="Branch", pem=_pem(ca[0])), request=_req(), current_user=manager, db=db)
    row = R.add_ca(R.CaIn(name="Branch", pem=_pem(ca[0])), request=_req(), current_user=admin, db=db)
    with pytest.raises(HTTPException) as e:
        R.add_ca(R.CaIn(name="Again", pem=_pem(ca[0])), request=_req(), current_user=admin, db=db)
    assert e.value.status_code == 409
    R.delete_ca(row["id"], request=_req(), current_user=admin, db=db)

    auto = TlsEndpoint(asset_id=a.id, host=a.ip_address, port=443, source=SOURCE_AUTO, first_seen_at=NOW)
    db.add(auto)
    db.flush()
    with pytest.raises(HTTPException) as e:
        R.delete_service(auto.id, request=_req(), current_user=admin, db=db)
    assert e.value.status_code == 400
    R.delete_service(out["id"], request=_req(), current_user=admin, db=db)

    R.scan(R.ScanIn(), request=_req(), current_user=manager, db=db)
    assert S.get_state(db)["requested"] is True
    actions = {r.action for r in db.query(SecurityAuditLog).filter(SecurityAuditLog.action.like("tls.%"))}
    assert {"tls.service_add", "tls.service_toggle", "tls.settings", "tls.ca_add", "tls.ca_delete",
            "tls.service_delete", "tls.scan"} <= actions


# ── remediation, risk, alerts, the asset page, the report ────────────────

def _service(db, asset, port=443, **kw):
    now = datetime.utcnow()
    fields = dict(fingerprint="cd" * 32, subject="CN=web.corp.local", issuer="CN=Corp Root CA",
                  names=["web.corp.local"], not_before=now - timedelta(days=100), not_after=now + timedelta(days=200),
                  trusted=True, self_signed=False, key_type="EC", key_bits=256, signature="sha256",
                  old_protocols=[], protocol="TLSv1.3", status="read", last_ok_at=now, last_scan_at=now)
    fields.update(kw)
    ep = TlsEndpoint(asset_id=asset.id, host=asset.ip_address, port=port, source=SOURCE_AUTO, first_seen_at=now,
                     **fields)
    db.add(ep)
    db.flush()
    return ep


def test_remediation_items_open_and_close(db):
    from app.models.remediation import RemediationItem
    from app.modules.remediation import service as rem
    a = _asset(db, ip="10.70.4.1")
    ep = _service(db, a, not_after=datetime.utcnow() - timedelta(days=3), trusted=False, self_signed=True,
                  old_protocols=["TLSv1"])
    rem.sync(db)
    items = {i.ref: i for i in db.query(RemediationItem).filter_by(asset_id=a.id, source="tls")}
    assert set(items) == {"expired", "untrusted", "old_protocol"}
    assert items["expired"].severity == "high" and items["untrusted"].severity == "medium"
    assert items["untrusted"].title == "The certificate of 10.70.4.1:443 is self-signed"
    assert items["old_protocol"].title == "10.70.4.1:443 still accepts TLS 1.0"
    assert items["expired"].title.startswith("The certificate of 10.70.4.1:443 expired on 20")
    ep.not_after, ep.trusted, ep.self_signed, ep.old_protocols = datetime.utcnow() + timedelta(days=365), True, False, []
    db.flush()
    rem.sync(db)
    for i in items.values():
        db.refresh(i)
        assert i.status == "resolved" and i.closed_reason == "fixed"


def test_risk_counts_certificate_problems(db):
    from app.modules.risk.service import AssetRiskCalculationService
    a = _asset(db, ip="10.70.5.1")
    _service(db, a, key_type="RSA", key_bits=1024)
    cv = AssetRiskCalculationService()._vulnerability_score(db, a.id)
    assert cv["score"] >= 30.0 and cv["certificate_problems"] == 1
    b = _asset(db, ip="10.70.5.2")
    _service(db, b, not_after=datetime.utcnow() - timedelta(days=1))
    assert AssetRiskCalculationService()._vulnerability_score(db, b.id)["score"] >= 60.0


def test_alerts(db):
    from app.modules.alerts.events import EVENTS
    now = datetime.utcnow()
    a = _asset(db, ip="10.70.6.1")
    gone = _service(db, a, not_after=now - timedelta(days=2))
    b = _asset(db, ip="10.70.6.2")
    _service(db, b, port=8443, not_after=now + timedelta(days=20))
    expired = EVENTS["tls.certificate_expired"].evaluate(db, {}, now, now)
    assert any(p.asset_id == a.id and p.key == f"tls:{gone.id}" and "10.70.6.1:443" in p.detail for p in expired)
    assert not any(p.asset_id == b.id for p in expired)
    soon = EVENTS["tls.certificate_expiring"].evaluate(db, {"days": 30}, now, now)
    assert any(p.asset_id == b.id for p in soon) and not any(p.asset_id == a.id for p in soon)
    assert not any(p.asset_id == b.id for p in EVENTS["tls.certificate_expiring"].evaluate(db, {"days": 10}, now, now))


def test_asset_page_lists_the_services(db):
    from app.modules.assets import overview360
    a = _asset(db, ip="10.70.7.1")
    _service(db, a, signature="sha1")
    out = overview360.build(db, a, _user(db))
    assert out["certificates"]["services"][0]["where"] == "10.70.7.1:443"
    assert [p["code"] for p in out["certificates"]["services"][0]["problems"]] == ["weak_signature"]


def test_the_tls_report_builds(db):
    import io

    from openpyxl import load_workbook

    from app.modules.reports import service as reports
    user = _user(db)
    a = _asset(db, ip="10.70.8.1")
    _service(db, a, not_after=datetime.utcnow() - timedelta(days=4))
    _service(db, a, port=636, old_protocols=["TLSv1", "TLSv1.1"], subject="CN=dc.corp.local")
    params = {"period": {"preset": "custom", "from": "2026-09-01", "to": "2026-09-30"}, "compare": False,
              "scope": {"mode": "assets", "values": [a.id]},
              "sections": ["overview", "expired", "expiring", "problems", "all"]}
    for lang in ("fa", "en"):
        r = reports.create(db, user, "tls_certificates", "", lang, ["pdf", "xlsx"], "internal", params)
        out = reports.build(db, r, now=datetime.utcnow())
        assert out["files"]["pdf"][1].startswith(b"%PDF")
        wb = load_workbook(io.BytesIO(out["files"]["xlsx"][1]))
        text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
        assert "10.70.8.1:443" in text and "TLS 1.0, TLS 1.1" in text and a.asset_name in text
