"""
End of life (app/modules/eol): the built-in catalog, recognising products and release cycles from the
software inventory, the status of each cycle on a day, the asset's typed operating system and entries made
by hand, the organisation's own products, updating the catalog online and from a file, the API (who may
change what), and what follows: remediation items, the risk score, alerts, the audit trail.

Needs a migrated PostgreSQL database; every test runs in a transaction that is rolled back.
"""
import json
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.core.database import engine
from app.models import Asset, AssetType, User, UserRole
from app.models.eol import SOURCE_CUSTOM, SOURCE_IMPORT, EolProduct
from app.models.remediation import RemediationItem
from app.models.security_audit_log import AuditLog as SecurityAuditLog
from app.models.software import SoftwareItem
from app.modules.eol import catalog as C
from app.modules.eol import match as M
from app.modules.eol import router as R
from app.modules.eol import service as S

TODAY = date(2026, 10, 18)


@pytest.fixture
def db():
    connection = engine.connect()
    trans = connection.begin()
    session = Session(bind=connection, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        session.close()
        trans.rollback()
        connection.close()


_seq = [0]


def _n():
    _seq[0] += 1
    return _seq[0]


def _user(db, role=UserRole.ADMIN):
    u = User(username=f"eol_{_n()}", hashed_password="x", role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _asset(db, os_name=None, os_version=None):
    t = AssetType(type_name=f"Server {_n()}", category="server")
    db.add(t)
    db.flush()
    a = Asset(asset_name=f"eol-{_n()}", asset_type_id=t.id, ip_address=f"10.68.0.{_n() % 250}",
              os_name=os_name, os_version=os_version)
    db.add(a)
    db.flush()
    return a


def _item(db, asset, name, version, kind="os", collector="linux"):
    db.add(SoftwareItem(asset_id=asset.id, collector=collector, kind=kind, name=name, version=version,
                        source="firmware" if kind == "firmware" else "distro"))
    db.flush()


def _req():
    return SimpleNamespace(client=SimpleNamespace(host="10.0.0.9"))


def _catalog():
    return {p["key"]: p for p in C.builtin()["products"]}


# ── catalog and matching ─────────────────────────────────────────────────

def test_builtin_catalog_is_valid_and_covers_the_recognised_products():
    data = C.builtin()
    assert data["source"].startswith("endoflife.date") and data["generated"]
    keys = {p["key"] for p in data["products"]}
    assert {k for k, _ in M.PRODUCT_NAMES} <= keys
    for p in data["products"]:
        assert p["cycles"], p["key"]
        assert all(C._valid_cycle(c) for c in p["cycles"]), p["key"]


@pytest.mark.parametrize("name,version,expected", [
    ("Ubuntu", "22.04", ("ubuntu", "22.04")),
    ("Red Hat Enterprise Linux", "9.4", ("rhel", "9")),
    ("CentOS Stream", "9", ("centos-stream", "9")),
    ("Windows Server 2019 Standard", "10.0.17763.5329", ("windows-server", "2019")),
    ("Windows Server 2012 R2 Datacenter", "6.3.9600.1", ("windows-server", "2012-r2")),
    ("Windows Server Datacenter", "10.0.20348.2", ("windows-server", "2022")),
    ("Windows 10 Enterprise LTSC 2021", "10.0.19044.1", ("windows", "10-21h2-e-lts")),
    ("Windows 11 Pro", "10.0.22631.4", ("windows", "11-23h2-w")),
    ("Fortinet FortiOS", "7.0.12", ("fortios", "7.0")),
    ("Cisco IOS XE", "17.03.04a", ("cisco-ios-xe", "17.3")),
    ("VMware ESXi", "7.0.3", ("vmware-esxi", "7.0")),
    ("Microsoft SQL Server 2016", "13.0.6300.2", ("mssqlserver", "13.0-sp3")),
    ("Microsoft SQL Server 2014", "12.0.5000.0", ("mssqlserver", "12.0-sp2")),
    ("Kubernetes", "v1.28.4", ("kubernetes", "1.28")),
    ("Docker Engine", "24.0.7", ("docker-engine", "24.0")),
    ("Ubuntu 18.04", None, ("ubuntu", "18.04")),
    ("Sophos Firewall OS", "20.0.1", (None, None)),
])
def test_products_and_cycles_are_recognised(name, version, expected):
    assert M.identify(name, version, _catalog()) == expected


def test_status_of_a_cycle_on_a_day():
    def st(**c):
        return S.status_of(c, TODAY, 180)["status"]
    assert st(security_until="2025-09-30") == "unsupported"
    assert st(security_until=True) == "unsupported"
    assert st(security_until="2023-05-31", extended_until="2028-04-01") == "extended"
    assert st(security_until="2023-05-31", extended_until="2026-04-23") == "unsupported"
    assert st(security_until="2027-01-31") == "ending"
    assert st(security_until="2031-10-14") == "supported"
    assert st(security_until=None) == "supported"
    assert S.status_of(None, TODAY, 180)["status"] == "unknown"
    assert S.status_of({"security_until": "2030-01-01", "active_until": "2025-01-01"}, TODAY, 180)["security_only"]


# ── the components of the assets ─────────────────────────────────────────

def test_components_from_inventory_typed_os_and_entries(db):
    linux, fw, typed, cam = _asset(db), _asset(db), _asset(db, "Windows Server 2012 R2", None), _asset(db)
    _item(db, linux, "Ubuntu", "18.04")                                   # extended (ESM) until 2028
    _item(db, linux, "Apache HTTP Server", "2.4.52", kind="service")      # supported
    _item(db, fw, "Fortinet FortiOS", "7.0.15", kind="firmware", collector="fortinet")   # ended 2025-09-30
    _item(db, fw, "Sophos Firewall OS", "20.0.1", kind="firmware", collector="fortinet")  # not in the catalog
    R.add_entry(R.EntryIn(asset_id=cam.id, name="Camera firmware", version="5.7",
                          support_until=TODAY + timedelta(days=30)), request=_req(), current_user=_user(db), db=db)
    comps = S.by_asset(S.components(db, [linux.id, fw.id, typed.id, cam.id], today=TODAY))
    status = {aid: {c.name: c.status for c in cs} for aid, cs in comps.items()}
    assert status[linux.id] == {"Ubuntu": "extended", "Apache HTTP Server": "supported"}
    assert status[fw.id] == {"Fortinet FortiOS": "unsupported", "Sophos Firewall OS": "unknown"}
    assert status[typed.id] == {"Windows Server 2012 R2": "unsupported"}         # ESU ended 2026-10-13
    assert status[cam.id] == {"Camera firmware": "ending"}
    assert S.worst(comps[fw.id]) == "unsupported" and S.worst(comps[linux.id]) == "extended"
    # an OS typed on an asset that has an inventory OS is not counted twice
    both = _asset(db, "Ubuntu", "16.04")
    _item(db, both, "Ubuntu", "24.04")
    assert [c.cycle for c in S.components(db, [both.id], today=TODAY)] == ["24.04"]


def test_the_organisations_own_product(db):
    admin = _user(db)
    fw = _asset(db)
    _item(db, fw, "Sophos Firewall OS", "19.5.3", kind="firmware")
    body = R.CustomIn(name="Sophos Firewall OS", match_names=["Sophos Firewall OS"],
                      cycles=[R.CycleIn(cycle="19.5", security_until=date(2025, 3, 31)),
                              R.CycleIn(cycle="20.0", security_until=date(2027, 3, 31))])
    key = R.add_custom(body, request=_req(), current_user=admin, db=db)["key"]
    comps = S.components(db, [fw.id], today=TODAY)
    assert [(c.product, c.cycle, c.status) for c in comps] == [(key, "19.5", "unsupported")]
    with pytest.raises(HTTPException) as e:
        R.add_custom(body, request=_req(), current_user=admin, db=db)
    assert e.value.status_code == 409
    manager = _user(db, UserRole.MANAGER)
    with pytest.raises(HTTPException) as e:
        R.add_custom(R.CustomIn(name="Other", match_names=["x"], cycles=body.cycles), request=_req(),
                     current_user=manager, db=db)
    assert e.value.status_code == 403


# ── updating the catalog ─────────────────────────────────────────────────

API_ANSWER = {"result": {"name": "fortios", "label": "FortiOS", "releases": [
    {"name": "7.0", "label": "7.0", "releaseDate": "2021-03-30", "isEol": True, "eolFrom": "2025-09-30",
     "isEoas": True, "eoasFrom": "2024-03-30", "isLts": False, "latest": {"name": "7.0.17"}},
    {"name": "7.6", "label": "7.6", "releaseDate": "2024-07-25", "isEol": False, "eolFrom": "2099-01-25",
     "isLts": False, "latest": {"name": "7.6.4"}}]}}


def test_files_and_api_answers_are_read(db):
    products = C.parse_file(json.dumps(API_ANSWER).encode())
    assert products[0]["key"] == "fortios"
    assert products[0]["cycles"][1]["security_until"] == "2099-01-25"
    own = C.parse_file(json.dumps({"products": [{"key": "ubuntu", "name": "Ubuntu", "cycles": [
        {"cycle": "22.04", "security_until": "2099-04-01"}]}]}).encode())
    assert own[0]["cycles"][0]["cycle"] == "22.04"
    for bad in (b"not json", json.dumps({"x": 1}).encode(),
                json.dumps({"products": [{"key": "a", "cycles": [{"cycle": "1", "security_until": "soon"}]}]}).encode()):
        with pytest.raises(ValueError):
            C.parse_file(bad)
    assert C.store(db, products, SOURCE_IMPORT, None) == 1
    loaded, _ = C.load(db)
    assert loaded["fortios"]["source"] == SOURCE_IMPORT and len(loaded["fortios"]["cycles"]) == 2
    # an organisation's own product is never overwritten by an import
    db.add(EolProduct(key="custom-x", name="X", source=SOURCE_CUSTOM, match_names=["X"],
                      cycles=[{"cycle": "1", "security_until": "2030-01-01"}]))
    db.flush()
    assert C.store(db, [{"key": "custom-x", "name": "X", "cycles": [{"cycle": "2"}]}], SOURCE_IMPORT, None) == 0


def test_online_update_takes_what_answers(db, monkeypatch):
    class Resp:
        def __init__(self, data, ok=True):
            self.data, self.ok = data, ok

        def raise_for_status(self):
            if not self.ok:
                raise RuntimeError("404")

        def json(self):
            return self.data

    class Sess:
        def get(self, url, **kw):
            return Resp(API_ANSWER) if url.endswith("/fortios") else Resp({}, ok=False)

    got, failed = C.fetch_online(Sess())
    assert [p["key"] for p in got] == ["fortios"] and "ubuntu" in failed
    monkeypatch.setattr(C, "fetch_online", lambda session=None: (got, failed))
    admin = _user(db)
    out = R.update_online(request=_req(), current_user=admin, db=db)
    assert out["updated"] == 1
    R.reset_product("fortios", request=_req(), current_user=admin, db=db)
    assert C.load(db)[0]["fortios"]["source"] == "builtin"


# ── API ──────────────────────────────────────────────────────────────────

def test_overview_asset_view_and_entries(db):
    user = _user(db)
    a = _asset(db)
    _item(db, a, "Fortinet FortiOS", "6.4.15", kind="firmware")       # ended 2024
    _item(db, a, "Windows Server 2025 Datacenter", "10.0.26100.1")      # supported to 2034
    over = R.overview(current_user=user, db=db)
    mine = [g for g in over["groups"] if any(x["id"] == a.id for x in g["assets"])]
    assert {(g["product"], g["cycle"], g["status"]) for g in mine} == {("fortios", "6.4", "unsupported"),
                                                                       ("windows-server", "2025", "supported")}
    view = R.asset_components(a.id, current_user=user, db=db)
    assert view["status"] == "unsupported" and view["components"][0]["product"] == "fortios"
    with pytest.raises(HTTPException):
        R.put_settings(R.SettingsIn(warn_days=90), request=_req(), current_user=_user(db, UserRole.MANAGER), db=db)
    assert R.put_settings(R.SettingsIn(warn_days=90), request=_req(), current_user=user, db=db)["warn_days"] == 90
    eid = R.add_entry(R.EntryIn(asset_id=a.id, name="BMC firmware", support_until=date(2020, 1, 1)),
                      request=_req(), current_user=user, db=db)["id"]
    R.delete_entry(eid, request=_req(), current_user=user, db=db)
    rows = db.query(SecurityAuditLog).filter(SecurityAuditLog.user_id == user.id).all()
    assert any(r.action == "eol.entry_add" and "BMC firmware" in r.detail for r in rows)


def test_routes_are_guarded():
    assert len(R.router.routes) >= 12
    for route in R.router.routes:
        deps = [d.call for d in route.dependant.dependencies]
        assert any(getattr(c, "permission", None) for c in deps), route.path


# ── remediation, risk, alerts ────────────────────────────────────────────

def test_remediation_items_open_and_close(db):
    from app.modules.remediation import service as rem
    a = _asset(db)
    _item(db, a, "Fortinet FortiOS", "6.4.15", kind="firmware")
    rem.sync(db)
    item = db.query(RemediationItem).filter_by(asset_id=a.id, source="eol").one()
    assert item.severity == "high" and item.ref == "fortios:6.4"
    assert item.title == "FortiOS 6.4 is past the end of support"
    db.query(SoftwareItem).filter_by(asset_id=a.id).update({"version": "7.6.4"})
    db.flush()
    rem.sync(db)
    db.refresh(item)
    assert item.status == "resolved" and item.closed_reason == "fixed"


def test_risk_counts_software_past_its_end(db):
    from app.modules.risk.service import AssetRiskCalculationService
    a = _asset(db)
    _item(db, a, "Fortinet FortiOS", "6.4.15", kind="firmware")
    cv = AssetRiskCalculationService()._vulnerability_score(db, a.id)
    assert cv["score"] >= 80.0 and cv["end_of_life"] == 1


def test_alerts(db):
    from app.modules.alerts.events import EVENTS
    a = _asset(db)
    _item(db, a, "Fortinet FortiOS", "6.4.15", kind="firmware")
    now = datetime.utcnow()
    problems = EVENTS["software.end_of_life"].evaluate(db, {}, now, now)
    assert any(p.asset_id == a.id and "FortiOS 6.4" in p.detail for p in problems)
    b = _asset(db)
    db.add(R.EolAssetEntry(asset_id=b.id, name="UPS card", support_until=now.date() + timedelta(days=20)))
    db.flush()
    soon = EVENTS["software.end_of_life_soon"].evaluate(db, {"days": 30}, now, now)
    assert any(p.asset_id == b.id for p in soon)
    assert not any(p.asset_id == b.id for p in EVENTS["software.end_of_life_soon"].evaluate(db, {"days": 10}, now, now))


def test_the_end_of_life_report_builds(db):
    import io

    from openpyxl import load_workbook

    from app.modules.reports import service as reports
    user = _user(db)
    a = _asset(db)
    _item(db, a, "Fortinet FortiOS", "6.4.15", kind="firmware")
    params = {"period": {"preset": "custom", "from": "2026-09-01", "to": "2026-09-30"}, "compare": False,
              "scope": {"mode": "assets", "values": [a.id]},
              "sections": ["overview", "unsupported", "extended", "ending", "unknown", "all"]}
    for lang in ("fa", "en"):
        r = reports.create(db, user, "end_of_life", "", lang, ["pdf", "xlsx"], "internal", params)
        out = reports.build(db, r, now=datetime.utcnow())
        assert out["files"]["pdf"][1].startswith(b"%PDF")
        wb = load_workbook(io.BytesIO(out["files"]["xlsx"][1]))
        text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)
        assert "FortiOS 6.4" in text and a.asset_name in text


@pytest.mark.parametrize("product_name, label, expected", [
    ("FortiOS", "7.0", "FortiOS 7.0"),
    ("Microsoft Windows Server", "Windows Server 2012 R2", "Windows Server 2012 R2"),
    ("Microsoft Windows", "11 24H2 IoT", "Microsoft Windows 11 24H2 IoT"),
    ("Amazon Linux", "AMI 2018.03", "Amazon Linux AMI 2018.03"),
])
def test_title_does_not_repeat_the_product(product_name, label, expected):
    c = S.Component(asset_id=1, name="x", version=None, kind="os", origin="inventory",
                    product="p", product_name=product_name, cycle="c", cycle_label=label)
    assert c.title == expected
    assert S.Component(asset_id=1, name="Sophos Firewall OS", version="19.5", kind="firmware",
                       origin="inventory").title == "Sophos Firewall OS"
